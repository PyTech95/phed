"""Iteration 9 backend tests for the 6-change PHED batch.

Covers:
  - Office documents Excel auto-approve import (POST /api/phed/office-import + sample + listing).
  - Mandatory documents on WATER_CONNECTION submit (blocks with 400).
  - Reopen path: after Submitted, another draft on same property must succeed;
    after Approved, draft must 409.
  - Attachment with GPS+captured_at fields; upload on Submitted OK, on Approved 409.
  - Map data returns phed_survey_status.
"""

import io
import os
import time
import uuid

import pytest
import requests
from openpyxl import Workbook, load_workbook

def _read_frontend_env():
    p = "/app/frontend/.env"
    if os.path.exists(p):
        for line in open(p):
            if line.startswith("REACT_APP_BACKEND_URL="):
                return line.split("=", 1)[1].strip()
    return None

BASE = (os.environ.get("REACT_APP_BACKEND_URL") or _read_frontend_env() or "").rstrip("/")
assert BASE, "REACT_APP_BACKEND_URL not configured"
TOWN = "THS"


def _login(username: str, password: str) -> str:
    r = requests.post(f"{BASE}/api/auth/login",
                      json={"username": username, "password": password},
                      headers={"X-Town-Code": TOWN}, timeout=30)
    assert r.status_code == 200, f"login {username}: {r.status_code} {r.text}"
    return r.json()["token"]


@pytest.fixture(scope="module")
def admin_headers():
    return {"Authorization": f"Bearer {_login('admin', 'PhedAdmin@2026')}", "X-Town-Code": TOWN}


@pytest.fixture(scope="module")
def surveyor_headers():
    return {"Authorization": f"Bearer {_login('surveyor1', 'Survey@2026')}", "X-Town-Code": TOWN}


@pytest.fixture(scope="module")
def surveyor_property(surveyor_headers):
    """Create a fresh field property owned by surveyor1."""
    body = {"owner_name": f"TEST Iter9 {uuid.uuid4().hex[:6]}", "mobile": "9999900001",
            "ward": "1", "colony": "Didar Nagar", "latitude": 29.97, "longitude": 76.83,
            "category": "Residential"}
    r = requests.post(f"{BASE}/api/phed/field-properties", json=body,
                      headers=surveyor_headers, timeout=30)
    assert r.status_code == 200, r.text
    return r.json()  # {id, property_id, property}


# ----------------------------- office import -----------------------------
class TestOfficeImport:
    def test_sample_endpoint_admin_only(self, admin_headers, surveyor_headers):
        r = requests.get(f"{BASE}/api/phed/office-import/sample", headers=admin_headers, timeout=30)
        assert r.status_code == 200
        assert "spreadsheetml" in r.headers.get("content-type", "")
        wb = load_workbook(io.BytesIO(r.content))
        assert wb.active.max_row >= 2

        r2 = requests.get(f"{BASE}/api/phed/office-import/sample", headers=surveyor_headers, timeout=30)
        assert r2.status_code == 403

    def test_office_import_flow(self, admin_headers, surveyor_headers, surveyor_property):
        pid_known = surveyor_property["property_id"]
        pid_unknown = f"NOPE-{uuid.uuid4().hex[:6].upper()}"

        # Row 3: unlinked consumer_id 4326479 (per spec — expected not-linked-to-property)
        wb = Workbook()
        ws = wb.active
        ws.append(["Property ID (PID)", "Consumer ID", "Owner Name", "Mobile", "Remarks"])
        ws.append([pid_known, "", "Iter9 Owner", "9876500000", "docs received"])
        ws.append([pid_unknown, "", "Ghost", "9876500001", "no such pid"])
        ws.append(["", "4326479", "Unlinked", "7206864297", "orphan consumer id"])
        buf = io.BytesIO(); wb.save(buf); buf.seek(0)

        files = {"file": ("iter9.xlsx", buf.getvalue(),
                         "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        r = requests.post(f"{BASE}/api/phed/office-import", files=files,
                          headers=admin_headers, timeout=60)
        assert r.status_code == 200, r.text
        data = r.json()
        assert data["total_rows"] == 3
        assert data["approved"] == 1, data
        assert data["not_found"] == 2, data
        assert len(data["unmatched"]) == 2
        reasons = [u["reason"] for u in data["unmatched"]]
        assert any("Property not found" in x or "consumer" in x.lower() for x in reasons)

        # property status flipped
        r2 = requests.get(f"{BASE}/api/phed/property/{surveyor_property['id']}",
                         headers=surveyor_headers, timeout=30)
        assert r2.status_code == 200
        assert r2.json()["property"]["phed_survey_status"] == "Approved"

        # Re-upload same → already_approved=1
        buf.seek(0)
        files = {"file": ("iter9.xlsx", buf.getvalue(),
                         "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        r3 = requests.post(f"{BASE}/api/phed/office-import", files=files,
                           headers=admin_headers, timeout=60)
        assert r3.status_code == 200
        d3 = r3.json()
        assert d3["already_approved"] == 1, d3
        assert d3["approved"] == 0, d3

        # office-imports listing exposes the run
        r4 = requests.get(f"{BASE}/api/phed/office-imports", headers=admin_headers, timeout=30)
        assert r4.status_code == 200
        rows = r4.json()
        assert any(row["filename"] == "iter9.xlsx" for row in rows)

    def test_non_admin_forbidden(self, surveyor_headers):
        wb = Workbook(); wb.active.append(["Property ID (PID)"]); wb.active.append(["X"])
        buf = io.BytesIO(); wb.save(buf); buf.seek(0)
        r = requests.post(f"{BASE}/api/phed/office-import",
                          files={"file": ("x.xlsx", buf.getvalue(),
                                          "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
                          headers=surveyor_headers, timeout=30)
        assert r.status_code == 403


# ----------------------------- reopen + mandatory docs -----------------------------
class TestReopenAndMandatoryDocs:
    @pytest.fixture(scope="class")
    def prop(self, surveyor_headers):
        body = {"owner_name": f"TEST Reopen {uuid.uuid4().hex[:6]}", "mobile": "9999900002",
                "ward": "1", "colony": "Didar Nagar", "latitude": 29.971, "longitude": 76.831,
                "category": "Residential"}
        r = requests.post(f"{BASE}/api/phed/field-properties", json=body,
                          headers=surveyor_headers, timeout=30)
        assert r.status_code == 200
        return r.json()

    def test_full_flow(self, surveyor_headers, admin_headers, prop):
        prid = prop["id"]

        # 1. draft with document_pending=True
        payload = {"property_record_id": prid, "survey_type": "WATER_CONNECTION",
                   "latitude": 29.971, "longitude": 76.831,
                   "water": {"has_connection": True, "document_pending": True,
                             "missing_documents": ["APPLICATION"]}}
        r = requests.post(f"{BASE}/api/phed/surveys/draft", json=payload,
                          headers=surveyor_headers, timeout=30)
        assert r.status_code == 200, r.text
        survey_id = r.json()["id"]

        # 2. submit → 400 with reason mentioning documents missing
        r2 = requests.post(f"{BASE}/api/phed/surveys/{survey_id}/submit",
                           headers=surveyor_headers, timeout=30)
        assert r2.status_code == 400, r2.text
        assert "document" in r2.text.lower() or "missing" in r2.text.lower()

        # 3. update draft with document_pending=False → submit → 200 Submitted
        payload["water"]["document_pending"] = False
        payload["water"]["missing_documents"] = []
        r3 = requests.post(f"{BASE}/api/phed/surveys/draft", json=payload,
                           headers=surveyor_headers, timeout=30)
        assert r3.status_code == 200

        r4 = requests.post(f"{BASE}/api/phed/surveys/{survey_id}/submit",
                           headers=surveyor_headers, timeout=30)
        assert r4.status_code == 200, r4.text
        assert r4.json()["status"] == "Submitted"

        # 4. Reopen: draft again on Submitted → must succeed
        r5 = requests.post(f"{BASE}/api/phed/surveys/draft", json=payload,
                           headers=surveyor_headers, timeout=30)
        assert r5.status_code == 200, f"reopen draft blocked: {r5.status_code} {r5.text}"

        # 5. resubmit → 200
        r6 = requests.post(f"{BASE}/api/phed/surveys/{survey_id}/submit",
                           headers=surveyor_headers, timeout=30)
        assert r6.status_code == 200, r6.text

        # 6. Admin approve
        r7 = requests.post(f"{BASE}/api/phed/surveys/{survey_id}/approve",
                           headers=admin_headers, timeout=30)
        assert r7.status_code == 200

        # 7. Now surveyor draft → 409
        r8 = requests.post(f"{BASE}/api/phed/surveys/draft", json=payload,
                           headers=surveyor_headers, timeout=30)
        assert r8.status_code == 409, r8.text


# ----------------------------- attachments GPS -----------------------------
def _tiny_jpeg() -> bytes:
    # 1x1 valid jpeg
    return bytes.fromhex(
        "ffd8ffe000104a46494600010101006000600000ffdb00430008060607060508070707"
        "0909080a0c140d0c0b0b0c1912130f141d1a1f1e1d1a1c1c20242e2720222c231c1c28"
        "37292c30313434341f27393d38323c2e333432ffc0000b080001000101011100ffc4001f"
        "0000010501010101010100000000000000000102030405060708090a0bffc400b510000"
        "20103030204030505040400000177000102031104051221314106135161072271143281"
        "91a1082342b1c11552d1f0243362727282090a161718191a25262728292a3435363738"
        "393a434445464748494a535455565758595a636465666768696a737475767778797a83"
        "8485868788898a92939495969798999aa2a3a4a5a6a7a8a9aab2b3b4b5b6b7b8b9bac2"
        "c3c4c5c6c7c8c9cad2d3d4d5d6d7d8d9dae1e2e3e4e5e6e7e8e9eaf1f2f3f4f5f6f7f8"
        "f9faffda0008010100003f00fbd0ffd9"
    )


class TestAttachmentsGps:
    def test_house_photo_with_gps(self, surveyor_headers, admin_headers):
        # fresh property
        body = {"owner_name": f"TEST Att {uuid.uuid4().hex[:6]}", "mobile": "9999900003",
                "ward": "1", "colony": "Didar Nagar", "latitude": 29.972, "longitude": 76.832,
                "category": "Residential"}
        pr = requests.post(f"{BASE}/api/phed/field-properties", json=body,
                           headers=surveyor_headers, timeout=30).json()
        prid = pr["id"]

        # draft (docs not pending)
        payload = {"property_record_id": prid, "survey_type": "WATER_CONNECTION",
                   "latitude": 29.972, "longitude": 76.832,
                   "water": {"has_connection": True, "document_pending": False,
                             "missing_documents": []}}
        d = requests.post(f"{BASE}/api/phed/surveys/draft", json=payload,
                          headers=surveyor_headers, timeout=30).json()
        sid = d["id"]

        # attachment with GPS
        files = {"file": ("house.jpg", _tiny_jpeg(), "image/jpeg")}
        form = {"attachment_type": "HOUSE_PHOTO", "latitude": "29.972",
                "longitude": "76.832", "captured_at": "2026-01-15T10:00:00Z",
                "gps_source": "device"}
        r = requests.post(f"{BASE}/api/phed/surveys/{sid}/attachments",
                          data=form, files=files, headers=surveyor_headers, timeout=30)
        assert r.status_code == 200, r.text
        att = r.json()
        assert att["latitude"] == 29.972
        assert att["longitude"] == 76.832
        assert att["captured_at"] == "2026-01-15T10:00:00Z"
        assert att["gps_source"] == "device"

        # GET property shows attachment fields
        p = requests.get(f"{BASE}/api/phed/property/{prid}",
                        headers=surveyor_headers, timeout=30).json()
        atts = p.get("survey", {}).get("attachments", [])
        assert any(a.get("latitude") == 29.972 and a.get("captured_at") for a in atts), atts

        # Submit → attach on Submitted still OK
        s = requests.post(f"{BASE}/api/phed/surveys/{sid}/submit",
                          headers=surveyor_headers, timeout=30)
        assert s.status_code == 200
        r2 = requests.post(f"{BASE}/api/phed/surveys/{sid}/attachments",
                           data=form, files={"file": ("h2.jpg", _tiny_jpeg(), "image/jpeg")},
                           headers=surveyor_headers, timeout=30)
        assert r2.status_code == 200, r2.text

        # Approve → attach → 409
        ap = requests.post(f"{BASE}/api/phed/surveys/{sid}/approve",
                          headers=admin_headers, timeout=30)
        assert ap.status_code == 200
        r3 = requests.post(f"{BASE}/api/phed/surveys/{sid}/attachments",
                           data=form, files={"file": ("h3.jpg", _tiny_jpeg(), "image/jpeg")},
                           headers=surveyor_headers, timeout=30)
        assert r3.status_code == 409


# ----------------------------- map data -----------------------------
class TestMapEndpoint:
    def test_employee_map_returns_phed_status(self, surveyor_headers):
        r = requests.get(f"{BASE}/api/map/employee-properties",
                        headers=surveyor_headers, timeout=60)
        assert r.status_code == 200
        js = r.json()
        # Endpoint may return list or dict with 'properties'
        props = js if isinstance(js, list) else js.get("properties") or js.get("items") or []
        assert props, "expected some properties"
        # at least one should carry phed_survey_status field (may be null)
        assert any("phed_survey_status" in p for p in props), \
            f"phed_survey_status missing from map payload; keys: {list(props[0].keys())[:20]}"
