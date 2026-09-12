"""PHED Water/Wastewater/Sewerage survey module: wards, consumers, connections, import, surveys, dashboard, export."""
import asyncio
import csv
import io
import logging
import math
import re
import uuid
from datetime import datetime, timezone
from typing import Any, Dict, List, Optional

from fastapi import APIRouter, Depends, File, Form, Header, HTTPException, Query, UploadFile
from fastapi.responses import Response, StreamingResponse
from openpyxl import Workbook, load_workbook
from pydantic import BaseModel, Field

from server import (
    ADMIN_VIEW_ROLES,
    clear_map_cache,
    get_current_user,
    get_db,
    get_fs,
    record_audit,
    save_file_to_gridfs,
)
from bson import ObjectId
from pymongo import ReturnDocument
import ward_master

logger = logging.getLogger("phed")
phed_router = APIRouter(prefix="/phed")

SERVICES = ["Water", "Sewer"]
CATEGORIES = ["Domestic", "Commercial", "Domestic-SC", "Other"]
SURVEY_TYPES = ["EXISTING_LINKED", "NEW_UNLISTED", "NO_CONNECTION", "WATER_CONNECTION"]
SURVEY_STATUSES = ["Not Started", "Draft", "Submitted", "Requires Review", "Approved", "Rejected", "Document Pending"]
OPEN_STATUSES = ["Submitted", "Requires Review", "Approved", "Document Pending"]  # counted as "surveyed" for pending math
DAILY_TARGET_DEFAULT = 30  # default surveys/day target per surveyor (leaderboard)
ATTACHMENT_TYPES = ["AADHAAR", "AADHAAR_FRONT", "AADHAAR_BACK", "REGISTRY", "BILL", "PROPERTY_FRONT", "OTHER", "APPLICATION", "PROPERTY_PROOF", "HOUSE_PHOTO", "DEATH_CERTIFICATE"]
ALLOWED_MIME = {"image/jpeg", "image/png", "image/webp", "application/pdf"}
MAX_ATTACHMENT_BYTES = 20 * 1024 * 1024  # HD photos / multi-page registry scans
GPS_DRIFT_FLAG_METERS = 100

HEADER_ALIASES = {
    "consumer_name": ["consumer name", "name of consumer", "name"],
    "fh_name": ["f h name", "fh name", "f/h name", "father husband name", "father/husband name", "fathers name", "husband name"],
    "ppp_id": ["head of family in ppp", "head of family in ppp id", "ppp id", "ppp", "hof ppp id", "family id"],
    "address": ["address", "address of consumer"],
    "locality": ["locality", "colony", "area"],
    "phone": ["phone no", "phone number", "phone", "mobile", "mobile no", "mobile number", "contact"],
    "consumer_id": ["consumer id", "consumer no", "consumer number", "cid"],
    "water_conn": ["water connection no", "water connection number", "water conn no", "water connection"],
    "sewer_conn": ["sewer connection no", "sewer connection number", "sewer conn no", "sewer connection", "sewerage connection no"],
    "category": ["type of connection", "connection type", "type", "category"],
}
FIELD_LABELS = {
    "consumer_name": "Consumer Name", "fh_name": "F/H Name", "ppp_id": "Head of Family in PPP", "address": "Address",
    "locality": "Locality", "phone": "Phone No.", "consumer_id": "Consumer ID", "water_conn": "Water Connection No.",
    "sewer_conn": "Sewer Connection No.", "category": "Type of Connection",
}
REQUIRED_FIELDS = ["consumer_name", "consumer_id"]


# ---------------- helpers ----------------
def now_iso() -> str:
    return datetime.now(timezone.utc).isoformat()


def norm_header(h: Any) -> str:
    s = re.sub(r"[.,:;_\-/()]+", " ", str(h or "").lower())
    return re.sub(r"\s+", " ", s).strip()


def norm_id(v: Any) -> str:
    return re.sub(r"[\s\-/]+", "", cell_str(v)).upper()


def norm_text(v: Any) -> str:
    return re.sub(r"\s+", " ", cell_str(v).lower()).strip()


def cell_str(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, float):
        if math.isnan(v):
            return ""
        if v.is_integer():
            return str(int(v))
        return repr(v)
    if isinstance(v, int):
        return str(v)
    if isinstance(v, datetime):
        return v.isoformat()
    return re.sub(r"\s+", " ", str(v)).strip()


def norm_category(raw: str):
    r = norm_text(raw).replace(" ", "").replace("-", "")
    if not r:
        return "Other", False
    if r in ("domesticsc", "domsc", "dsc"):
        return "Domestic-SC", True
    if r.startswith("dom"):
        return "Domestic", True
    if r.startswith("com"):
        return "Commercial", True
    return "Other", False


def haversine_m(lat1, lon1, lat2, lon2) -> float:
    R = 6371000.0
    p1, p2 = math.radians(lat1), math.radians(lat2)
    dphi = p2 - p1
    dl = math.radians(lon2 - lon1)
    a = math.sin(dphi / 2) ** 2 + math.cos(p1) * math.cos(p2) * math.sin(dl / 2) ** 2
    return 2 * R * math.asin(math.sqrt(a))


def polygon_centroid(points):
    """points: list of [lat, lng]. Returns [lat, lng] centroid (area-weighted)."""
    pts = [p for p in (points or []) if isinstance(p, (list, tuple)) and len(p) == 2]
    n = len(pts)
    if n == 0:
        return None
    if n < 3:
        return [sum(p[0] for p in pts) / n, sum(p[1] for p in pts) / n]
    a = cx = cy = 0.0
    for i in range(n):
        y0, x0 = float(pts[i][0]), float(pts[i][1])
        y1, x1 = float(pts[(i + 1) % n][0]), float(pts[(i + 1) % n][1])
        cross = x0 * y1 - x1 * y0
        a += cross
        cx += (x0 + x1) * cross
        cy += (y0 + y1) * cross
    if abs(a) < 1e-12:
        return [sum(p[0] for p in pts) / n, sum(p[1] for p in pts) / n]
    a *= 0.5
    return [cy / (6 * a), cx / (6 * a)]  # [lat, lng]


def mask_phone(p: str) -> str:
    return p if len(p) < 6 else f"{p[:2]}{'*' * (len(p) - 4)}{p[-2:]}"


def is_admin(u): return u["role"] == "ADMIN"
def is_officer(u): return u["role"] in ADMIN_VIEW_ROLES
def require_admin(u):
    if not is_admin(u):
        raise HTTPException(403, "Administrator access required")
def require_officer(u):
    if not is_officer(u):
        raise HTTPException(403, "Officer/Administrator access required")


async def audit(actor: dict, action: str, entity_type: str, entity_id: str, before=None, after=None, meta=None):
    await get_db().phed_audit_logs.insert_one({
        "id": str(uuid.uuid4()), "actor_id": actor["id"], "actor_name": actor.get("name"), "actor_role": actor.get("role"),
        "action": action, "entity_type": entity_type, "entity_id": entity_id,
        "before": before, "after": after, "meta": meta, "timestamp": now_iso(),
    })


async def surveyor_can_access_property(user: dict, prop: dict) -> bool:
    if is_officer(user):
        return True
    return prop.get("assigned_employee_id") == user["id"] or user["id"] in (prop.get("assigned_employee_ids") or [])


async def get_property_or_403(property_record_id: str, user: dict) -> dict:
    prop = await get_db().properties.find_one({"id": property_record_id}, {"_id": 0})
    if not prop:
        raise HTTPException(404, "Property not found")
    if not await surveyor_can_access_property(user, prop):
        raise HTTPException(403, "Property not assigned to you")
    return prop


async def attach_connections(consumers: List[dict]) -> List[dict]:
    if not consumers:
        return consumers
    refs = [c["id"] for c in consumers]
    conns = await get_db().phed_connections.find({"consumer_ref": {"$in": refs}, "is_active": True}, {"_id": 0}).to_list(None)
    by_ref: Dict[str, list] = {}
    for c in conns:
        by_ref.setdefault(c["consumer_ref"], []).append(c)
    for c in consumers:
        c["connections"] = sorted(by_ref.get(c["id"], []), key=lambda x: (x["service"], x["connection_number"]))
        c["water_count"] = sum(1 for x in c["connections"] if x["service"] == "Water")
        c["sewer_count"] = sum(1 for x in c["connections"] if x["service"] == "Sewer")
        c["total_connections"] = len(c["connections"])
    return consumers


async def ensure_indexes(db):
    await db.phed_consumers.create_index("consumer_id_norm", background=True)
    await db.phed_consumers.create_index("phone_norm", background=True)
    await db.phed_consumers.create_index("name_norm", background=True)
    await db.phed_consumers.create_index("linked_property_id", background=True)
    await db.phed_consumers.create_index([("ward_id", 1), ("colony_name", 1)], background=True)
    await db.phed_consumers.create_index("status", background=True)
    await db.phed_connections.create_index([("service", 1), ("connection_number_norm", 1)], background=True)
    await db.phed_connections.create_index("consumer_ref", background=True)
    await db.phed_connections.create_index("linked_property_id", background=True)
    await db.phed_surveys.create_index([("property_record_id", 1), ("status", 1)], background=True)
    await db.phed_surveys.create_index("surveyor_id", background=True)
    await db.phed_surveys.create_index("status", background=True)
    await db.phed_surveys.create_index("submitted_at", background=True)
    await db.phed_audit_logs.create_index("entity_id", background=True)
    await db.phed_wards.create_index("ward_number", background=True)


# ---------------- wards ----------------
class WardIn(BaseModel):
    ward_number: str
    name: Optional[str] = None
    is_active: bool = True
    colonies: Optional[List[str]] = None


@phed_router.get("/wards")
async def list_wards(user: dict = Depends(get_current_user)):
    wards = await get_db().phed_wards.find({}, {"_id": 0}).sort("ward_number", 1).to_list(None)
    all_colonies = [c for c in await get_db().properties.distinct("colony") if c]
    mapped = {c for w in wards for c in (w.get("colonies") or [])}
    return {"wards": wards, "unmapped_colonies": sorted(c for c in all_colonies if c not in mapped)}


@phed_router.post("/wards")
async def create_ward(data: WardIn, user: dict = Depends(get_current_user)):
    require_admin(user)
    if await get_db().phed_wards.find_one({"ward_number": data.ward_number.strip()}):
        raise HTTPException(400, "Ward already exists")
    doc = {"id": str(uuid.uuid4()), "ward_number": data.ward_number.strip(), "name": (data.name or f"Ward {data.ward_number}").strip(),
           "is_active": data.is_active, "colonies": data.colonies or [], "created_at": now_iso(), "updated_at": now_iso()}
    await get_db().phed_wards.insert_one(doc)
    doc.pop("_id", None)
    await audit(user, "WARD_CREATE", "ward", doc["id"], after=doc)
    return doc


@phed_router.put("/wards/{ward_id}")
async def update_ward(ward_id: str, data: WardIn, user: dict = Depends(get_current_user)):
    require_admin(user)
    before = await get_db().phed_wards.find_one({"id": ward_id}, {"_id": 0})
    if not before:
        raise HTTPException(404, "Ward not found")
    upd = {"ward_number": data.ward_number.strip(), "name": (data.name or f"Ward {data.ward_number}").strip(),
           "is_active": data.is_active, "colonies": data.colonies if data.colonies is not None else before.get("colonies", []),
           "updated_at": now_iso()}
    await get_db().phed_wards.update_one({"id": ward_id}, {"$set": upd})
    await get_db().phed_consumers.update_many({"ward_id": ward_id}, {"$set": {"ward_number": upd["ward_number"]}})
    await audit(user, "WARD_UPDATE", "ward", ward_id, before=before, after={**before, **upd})
    return {**before, **upd}


@phed_router.delete("/wards/{ward_id}")
async def delete_ward(ward_id: str, user: dict = Depends(get_current_user)):
    require_admin(user)
    if await get_db().phed_consumers.count_documents({"ward_id": ward_id}) > 0:
        raise HTTPException(400, "Ward has consumers; deactivate it instead")
    await get_db().phed_wards.delete_one({"id": ward_id})
    await audit(user, "WARD_DELETE", "ward", ward_id)
    return {"message": "Ward deleted"}

class WardAssignIn(BaseModel):
    employee_id: str
    colony: Optional[str] = None      # restrict to one colony inside the ward
    only_unassigned: bool = True


@phed_router.post("/wards/{ward_id}/assign")
async def assign_ward(ward_id: str, data: WardAssignIn, user: dict = Depends(get_current_user)):
    """Assign PHED work by ward (all colonies) or by a single colony; reuses the existing property assignment fields."""
    require_admin(user)
    db = get_db()
    ward = await db.phed_wards.find_one({"id": ward_id}, {"_id": 0})
    if not ward:
        raise HTTPException(404, "Ward not found")
    from server import master_db
    emp = await master_db.users.find_one({"id": data.employee_id, "role": {"$in": ["SURVEYOR", "EMPLOYEE"]}}, {"_id": 0, "id": 1, "name": 1})
    if not emp:
        raise HTTPException(404, "Surveyor not found")
    colonies = [data.colony] if data.colony else ward.get("colonies", [])
    if data.colony and data.colony not in ward.get("colonies", []):
        raise HTTPException(400, "Colony does not belong to this ward")
    q: Dict[str, Any] = {"colony": {"$in": colonies}}
    if data.only_unassigned:
        q["$or"] = [{"assigned_employee_id": None}, {"assigned_employee_id": {"$exists": False}}]
    res = await db.properties.update_many(q, {"$set": {"assigned_employee_id": emp["id"], "assigned_employee_name": emp["name"], "assigned_at": now_iso()}})
    await clear_map_cache()
    await audit(user, "WARD_ASSIGN", "ward", ward_id, meta={"employee_id": emp["id"], "colonies": colonies, "properties": res.modified_count})
    await record_audit(user, "PHED_WARD_ASSIGN", "ward", ward_id, {"employee": emp["name"], "colonies": colonies, "properties_assigned": res.modified_count})
    return {"message": f"Assigned {res.modified_count} properties to {emp['name']}", "assigned": res.modified_count}




# ---------------- import ----------------
def parse_workbook(content: bytes):
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    header_row = None
    header_idx = 0
    for i, r in enumerate(rows):
        if r and sum(1 for c in r if cell_str(c)) >= 3:
            header_row = r
            header_idx = i
            break
    if header_row is None:
        raise HTTPException(400, "Could not find a header row in the workbook")
    mapping: Dict[str, int] = {}
    unmapped = []
    for col, h in enumerate(header_row):
        nh = norm_header(h)
        if not nh:
            continue
        hit = None
        for field, aliases in HEADER_ALIASES.items():
            if field in mapping:
                continue
            if nh in aliases or nh.replace(" ", "") in [a.replace(" ", "") for a in aliases]:
                hit = field
                break
        if hit:
            mapping[hit] = col
        else:
            unmapped.append(cell_str(h))
    data_rows = []
    for i, r in enumerate(rows):
        if r is None or not any(cell_str(c) for c in r):
            continue
        data_rows.append((header_idx + i + 2, r))
    return mapping, unmapped, data_rows, [cell_str(h) for h in header_row]


def row_to_record(mapping, row_tuple):
    row_no, r = row_tuple
    def g(f):
        idx = mapping.get(f)
        return cell_str(r[idx]) if idx is not None and idx < len(r) else ""
    cat, known = norm_category(g("category"))
    rec = {
        "row": row_no, "consumer_name": g("consumer_name"), "fh_name": g("fh_name"), "ppp_id": g("ppp_id"),
        "address": g("address"), "locality": g("locality"), "phone": g("phone"), "consumer_id": g("consumer_id"),
        "water_conn": g("water_conn"), "sewer_conn": g("sewer_conn"), "category_raw": g("category"), "category": cat,
        "category_known": known,
    }
    errors, warnings = [], []
    if not rec["consumer_id"]:
        errors.append("Consumer ID is required")
    if not rec["consumer_name"]:
        errors.append("Consumer Name is required")
    if not rec["water_conn"] and not rec["sewer_conn"]:
        warnings.append("No Water or Sewer connection number")
    for f, lbl in (("phone", "Phone"), ("fh_name", "F/H Name"), ("ppp_id", "PPP ID"), ("sewer_conn", "Sewer connection")):
        if not rec[f]:
            warnings.append(f"{lbl} missing")
    if rec["category_raw"] and not known:
        warnings.append(f"Unknown connection type '{rec['category_raw']}' stored as Other")
    rec["errors"], rec["warnings"] = errors, warnings
    return rec


@phed_router.get("/import/sample")
async def import_sample(user: dict = Depends(get_current_user)):
    """Downloadable .xlsx template with the exact headers the importer expects + example rows."""
    require_admin(user)
    headers = [
        "Consumer Name", "F/H Name", "Head of Family in PPP", "Address", "Locality",
        "Phone No.", "Consumer ID", "Water Connection No.", "Sewer Connection No.", "Type of Connection",
    ]
    examples = [
        ["RAMESH KUMAR", "SOM NATH", "2EFG7557", "GALI NO 4 DIDAR NAGAR", "DIDAR NAGAR", "9991626167", "3491531", "KU1w2366", "KU1s207", "Domestic"],
        ["BABITA RANI", "RANJEET SINGH", "", "NEW SHANTI NAGAR", "SHANTI NAGAR", "9466787688", "3529391", "KU1w2601", "", "Domestic"],
        ["SUNIL TRADERS", "", "", "MAIN MARKET DIDAR NAGAR", "DIDAR NAGAR", "9812345678", "3529400", "", "KU1s210", "Commercial"],
    ]
    wb = Workbook()
    ws = wb.active
    ws.title = "Consumers"
    ws.append(headers)
    for row in examples:
        ws.append(row)
    for i, h in enumerate(headers, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = max(14, len(h) + 2)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="PHED_import_sample.xlsx"'},
    )


@phed_router.post("/import/validate")
async def import_validate(
    file: UploadFile = File(...), ward_id: str = Form(...), default_colony: str = Form(""), mode: str = Form("import_and_update"),
    user: dict = Depends(get_current_user),
):
    require_admin(user)
    if mode not in ("validate_only", "new_only", "update_only", "import_and_update"):
        raise HTTPException(400, "Invalid import mode")
    ward = await get_db().phed_wards.find_one({"id": ward_id}, {"_id": 0})
    if not ward:
        raise HTTPException(400, "Select a valid ward")
    if not (file.filename or "").lower().endswith((".xlsx", ".xlsm")):
        raise HTTPException(400, "Only .xlsx files are supported")
    content = await file.read()
    mapping, unmapped, data_rows, raw_headers = parse_workbook(content)
    missing = [FIELD_LABELS[f] for f in REQUIRED_FIELDS if f not in mapping]
    if missing:
        raise HTTPException(400, f"Missing required column(s): {', '.join(missing)}. Found headers: {', '.join(raw_headers)}")

    records = [row_to_record(mapping, rt) for rt in data_rows]
    db = get_db()
    ids = list({norm_id(r["consumer_id"]) for r in records if r["consumer_id"]})
    existing = {c["consumer_id_norm"]: c for c in await db.phed_consumers.find({"consumer_id_norm": {"$in": ids}}, {"_id": 0, "id": 1, "consumer_id_norm": 1, "consumer_id": 1}).to_list(None)}
    conn_norms = list({norm_id(x) for r in records for x in (r["water_conn"], r["sewer_conn"]) if x})
    existing_conns = await db.phed_connections.find({"connection_number_norm": {"$in": conn_norms}, "is_active": True}, {"_id": 0, "connection_number_norm": 1, "service": 1, "consumer_id": 1}).to_list(None)
    # Connection numbers are unique PER CONSUMER (Consumer ID is the unique key); the same number may appear under different consumers.
    conn_owned = {(c["service"], c["connection_number_norm"], norm_id(c["consumer_id"])) for c in existing_conns}

    seen_ids: Dict[str, int] = {}
    seen_conns: set = set()
    seen_nums: Dict[tuple, set] = {}
    counts = dict(total_rows=len(records), valid_rows=0, invalid_rows=0, new_consumers=0, existing_consumers=0,
                  duplicate_consumer_ids=0, new_water_connections=0, new_sewer_connections=0, rows_with_missing_optional=0, conflicts=0,
                  shared_connection_numbers=0)
    conflicts = []
    for r in records:
        cid = norm_id(r["consumer_id"])
        if cid:
            if cid in seen_ids:
                counts["duplicate_consumer_ids"] += 1
                r["warnings"].append(f"Duplicate Consumer ID in file (first at row {seen_ids[cid]})")
            else:
                seen_ids[cid] = r["row"]
        if r["errors"]:
            counts["invalid_rows"] += 1
            continue
        counts["valid_rows"] += 1
        if r["warnings"]:
            counts["rows_with_missing_optional"] += 1
        r["exists"] = cid in existing
        counts["existing_consumers" if r["exists"] else "new_consumers"] += 1
        for svc, key, ck in (("Water", "water_conn", "new_water_connections"), ("Sewer", "sewer_conn", "new_sewer_connections")):
            if not r[key]:
                continue
            cn = norm_id(r[key])
            fk = (svc, cn, cid)
            if fk in seen_conns:
                counts["conflicts"] += 1
                conflicts.append({"row": r["row"], "service": svc, "connection_number": r[key], "consumer_id": r["consumer_id"], "existing_consumer_id": r["consumer_id"],
                                  "reason": "Same connection number repeated for this consumer in the file"})
                continue
            seen_conns.add(fk)
            holders = seen_nums.setdefault((svc, cn), set())
            if holders and cid not in holders:
                counts["shared_connection_numbers"] += 1
            holders.add(cid)
            if fk not in conn_owned:
                counts[ck] += 1

    import_id = str(uuid.uuid4())
    doc = {
        "id": import_id, "filename": re.sub(r"[^A-Za-z0-9._ -]", "_", file.filename or "upload.xlsx"), "ward_id": ward_id, "ward_number": ward["ward_number"],
        "default_colony": default_colony.strip(), "mode": mode, "uploaded_by": user["id"], "uploaded_by_name": user.get("name"),
        "status": "Validated", "header_mapping": {FIELD_LABELS[f]: raw_headers[i] for f, i in mapping.items()}, "unmapped_headers": unmapped,
        "counts": counts, "conflicts": conflicts, "errors": [{"row": r["row"], "reason": "; ".join(r["errors"])} for r in records if r["errors"]],
        "progress": {"processed": 0, "total": counts["valid_rows"]}, "result": None, "created_at": now_iso(), "completed_at": None,
    }
    await db.phed_imports.insert_one(doc)
    await db.phed_import_staging.insert_one({"import_id": import_id, "records": records})
    doc.pop("_id", None)
    preview = [{k: r[k] for k in ("row", "consumer_id", "consumer_name", "fh_name", "ppp_id", "address", "locality", "phone", "water_conn", "sewer_conn", "category", "warnings", "errors")} for r in records[:10]]
    await audit(user, "IMPORT_VALIDATE", "import", import_id, meta={"filename": doc["filename"], "counts": counts})
    return {**doc, "preview": preview}


async def _run_import(import_id: str, user: dict):
    db = get_db()
    imp = await db.phed_imports.find_one({"id": import_id}, {"_id": 0})
    staging = await db.phed_import_staging.find_one({"import_id": import_id}, {"_id": 0})
    records = [r for r in staging["records"] if not r["errors"]]
    mode = imp["mode"]
    result = dict(created_consumers=0, updated_consumers=0, skipped_consumers=0, created_connections=0, skipped_connections=0, conflicts=len(imp["conflicts"]))
    conflict_keys = {(c["row"], c["service"]) for c in imp["conflicts"]}
    processed = 0
    ts = now_iso()
    for r in records:
        try:
            cid_norm = norm_id(r["consumer_id"])
            existing = await db.phed_consumers.find_one({"consumer_id_norm": cid_norm}, {"_id": 0})
            colony = r["locality"] or imp["default_colony"]
            base = {
                "consumer_id": r["consumer_id"], "consumer_id_norm": cid_norm, "consumer_name": r["consumer_name"], "name_norm": norm_text(r["consumer_name"]),
                "fh_name": r["fh_name"], "ppp_id": r["ppp_id"], "address": r["address"], "address_norm": norm_text(r["address"]),
                "locality": r["locality"], "colony_name": colony, "phone": r["phone"], "phone_norm": norm_id(r["phone"]),
                "ward_id": imp["ward_id"], "ward_number": imp["ward_number"], "category": r["category"], "category_raw": r["category_raw"],
                "source": "imported", "source_file": imp["filename"], "source_row": r["row"], "import_id": import_id,
                "original_values": {FIELD_LABELS[k]: r[k] for k in ("consumer_name", "fh_name", "ppp_id", "address", "locality", "phone", "consumer_id", "water_conn", "sewer_conn")} | {"Type of Connection": r["category_raw"]},
                "updated_at": ts,
            }
            if existing:
                if mode == "new_only":
                    result["skipped_consumers"] += 1
                    consumer_ref = existing["id"]
                else:
                    keep = {k: existing[k] for k in ("linked_property_id", "linked_property_number", "status", "is_active", "created_at", "id") if k in existing}
                    await db.phed_consumers.update_one({"id": existing["id"]}, {"$set": {**base, **keep}})
                    result["updated_consumers"] += 1
                    consumer_ref = existing["id"]
            else:
                if mode == "update_only":
                    result["skipped_consumers"] += 1
                    processed += 1
                    continue
                consumer_ref = str(uuid.uuid4())
                await db.phed_consumers.insert_one({**base, "id": consumer_ref, "linked_property_id": None, "linked_property_number": None,
                                                    "status": "Imported", "is_active": True, "provisional_ref": None, "created_at": ts})
                result["created_consumers"] += 1
            for svc, key in (("Water", "water_conn"), ("Sewer", "sewer_conn")):
                if not r[key] or (r["row"], svc) in conflict_keys:
                    if r[key]:
                        result["skipped_connections"] += 1
                    continue
                cn = norm_id(r[key])
                if await db.phed_connections.find_one({"consumer_ref": consumer_ref, "service": svc, "connection_number_norm": cn, "is_active": True}, {"_id": 1}):
                    result["skipped_connections"] += 1
                    continue
                await db.phed_connections.insert_one({
                    "id": str(uuid.uuid4()), "consumer_ref": consumer_ref, "consumer_id": r["consumer_id"], "connection_number": r[key],
                    "connection_number_norm": cn, "service": svc, "category": r["category"], "category_raw": r["category_raw"], "source": "imported",
                    "status": "Active", "is_active": True, "linked_property_id": None, "import_id": import_id, "created_at": ts, "updated_at": ts,
                })
                result["created_connections"] += 1
        except Exception as e:  # keep going; record row failure
            logger.exception("import row failed")
            await db.phed_imports.update_one({"id": import_id}, {"$push": {"errors": {"row": r["row"], "reason": f"Processing error: {e}"}}})
        processed += 1
        if processed % 25 == 0:
            await db.phed_imports.update_one({"id": import_id}, {"$set": {"progress.processed": processed}})
    await db.phed_imports.update_one({"id": import_id}, {"$set": {"status": "Completed", "result": result, "progress.processed": processed, "completed_at": now_iso()}})
    await db.phed_import_staging.delete_one({"import_id": import_id})
    await audit(user, "IMPORT_COMMIT", "import", import_id, meta=result)


@phed_router.post("/import/{import_id}/commit")
async def import_commit(import_id: str, user: dict = Depends(get_current_user)):
    require_admin(user)
    imp = await get_db().phed_imports.find_one({"id": import_id}, {"_id": 0})
    if not imp:
        raise HTTPException(404, "Import not found")
    if imp["status"] != "Validated":
        raise HTTPException(409, f"Import is already {imp['status']}")
    if imp["mode"] == "validate_only":
        await get_db().phed_imports.update_one({"id": import_id}, {"$set": {"status": "Validated Only", "completed_at": now_iso()}})
        await get_db().phed_import_staging.delete_one({"import_id": import_id})
        return {"message": "Validation-only import closed", "status": "Validated Only"}
    await get_db().phed_imports.update_one({"id": import_id}, {"$set": {"status": "Processing"}})
    asyncio.create_task(_run_import(import_id, user))
    return {"message": "Import started", "status": "Processing"}


@phed_router.get("/imports")
async def list_imports(user: dict = Depends(get_current_user)):
    require_officer(user)
    items = await get_db().phed_imports.find({}, {"_id": 0, "errors": 0, "conflicts": 0}).sort("created_at", -1).limit(100).to_list(None)
    return {"imports": items}


@phed_router.get("/import/{import_id}")
async def get_import(import_id: str, user: dict = Depends(get_current_user)):
    require_officer(user)
    imp = await get_db().phed_imports.find_one({"id": import_id}, {"_id": 0})
    if not imp:
        raise HTTPException(404, "Import not found")
    return imp


@phed_router.get("/import/{import_id}/errors.csv")
async def import_errors_csv(import_id: str, user: dict = Depends(get_current_user)):
    require_officer(user)
    imp = await get_db().phed_imports.find_one({"id": import_id}, {"_id": 0})
    if not imp:
        raise HTTPException(404, "Import not found")
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["Row", "Type", "Reason", "Service", "Connection Number", "Consumer ID", "Existing Consumer ID"])
    for e in imp.get("errors", []):
        w.writerow([e["row"], "Error", e["reason"], "", "", "", ""])
    for c in imp.get("conflicts", []):
        w.writerow([c["row"], "Conflict", c["reason"], c["service"], c["connection_number"], c["consumer_id"], c["existing_consumer_id"]])
    return Response(buf.getvalue(), media_type="text/csv", headers={"Content-Disposition": f'attachment; filename="import_errors_{import_id[:8]}.csv"'})


# ---------------- consumers ----------------
class ConsumerIn(BaseModel):
    consumer_name: str
    fh_name: Optional[str] = ""
    ppp_id: Optional[str] = ""
    address: Optional[str] = ""
    locality: Optional[str] = ""
    phone: Optional[str] = ""
    consumer_id: Optional[str] = ""
    ward_id: Optional[str] = None
    category: Optional[str] = "Domestic"
    linked_property_id: Optional[str] = None
    connections: Optional[List[Dict[str, str]]] = None  # [{service, connection_number, category}]


class ConnectionIn(BaseModel):
    service: str
    connection_number: str
    category: Optional[str] = "Domestic"
    linked_property_id: Optional[str] = None
    remarks: Optional[str] = None


class LinkIn(BaseModel):
    property_record_id: str
    confirm: bool = False


def consumer_view(c: dict, user: dict) -> dict:
    if not is_officer(user) and c.get("phone"):
        c = {**c, "phone_masked": mask_phone(c["phone"])}
    return c


async def property_brief(property_record_id: Optional[str]):
    if not property_record_id:
        return None
    return await get_db().properties.find_one({"id": property_record_id}, {"_id": 0, "id": 1, "property_id": 1, "owner_name": 1, "address": 1, "colony": 1, "ward": 1, "latitude": 1, "longitude": 1, "status": 1, "polygon": 1, "polygon_centroid": 1})


@phed_router.get("/consumers/search")
async def search_consumers(q: str = Query(..., min_length=1), limit: int = 20, user: dict = Depends(get_current_user)):
    db = get_db()
    qn, qt = norm_id(q), norm_text(q)
    exact_ids = set()
    exact = await db.phed_consumers.find({"$or": [{"consumer_id_norm": qn}, {"phone_norm": qn}, {"provisional_ref": qn}]}, {"_id": 0}).limit(limit).to_list(None)
    exact_ids.update(c["id"] for c in exact)
    conn_hits = await db.phed_connections.find({"connection_number_norm": qn, "is_active": True}, {"_id": 0, "consumer_ref": 1}).to_list(None)
    refs = [c["consumer_ref"] for c in conn_hits if c["consumer_ref"] not in exact_ids]
    if refs:
        more = await db.phed_consumers.find({"id": {"$in": refs}}, {"_id": 0}).to_list(None)
        exact.extend(more)
        exact_ids.update(c["id"] for c in more)
    prop_hits = await db.properties.find({"property_id": {"$regex": f"^{re.escape(q.strip())}$", "$options": "i"}}, {"_id": 0, "id": 1}).limit(5).to_list(None)
    if prop_hits:
        more = await db.phed_consumers.find({"linked_property_id": {"$in": [p["id"] for p in prop_hits]}, "id": {"$nin": list(exact_ids)}}, {"_id": 0}).to_list(None)
        exact.extend(more)
        exact_ids.update(c["id"] for c in more)
    partial = []
    if len(exact) < limit and len(qt) >= 2:
        remaining = limit - len(exact)
        skip_ids = set(exact_ids)
        # 1) names / father-husband names that START WITH the query -> best suggestions
        prefix_rx = {"$regex": "^" + re.escape(qt), "$options": "i"}
        name_prefix = await db.phed_consumers.find(
            {"id": {"$nin": list(skip_ids)}, "$or": [{"name_norm": prefix_rx}, {"fh_name": prefix_rx}]},
            {"_id": 0}).limit(remaining).to_list(None)
        partial.extend(name_prefix)
        skip_ids.update(c["id"] for c in name_prefix)
        # 2) fill the rest with broader matches (name/father/address/locality/id/phone substrings)
        if len(partial) < remaining:
            rx = {"$regex": re.escape(qt).replace(r"\ ", r"\s*"), "$options": "i"}
            more = await db.phed_consumers.find(
                {"id": {"$nin": list(skip_ids)}, "$or": [{"name_norm": rx}, {"fh_name": rx}, {"address_norm": rx}, {"locality": rx}, {"consumer_id_norm": {"$regex": re.escape(qn), "$options": "i"}}, {"phone_norm": {"$regex": re.escape(qn)}}]},
                {"_id": 0}).limit(remaining - len(partial)).to_list(None)
            partial.extend(more)
    results = await attach_connections(exact + partial)
    for c in results:
        c["match"] = "exact" if c["id"] in exact_ids else "partial"
        c["linked_property"] = await property_brief(c.get("linked_property_id"))

    def _rank(c):
        # exact hits first, then name prefix, then word-prefix, then substring, then other-field matches
        if c["id"] in exact_ids:
            return (0, 0, c.get("name_norm") or "")
        nm = c.get("name_norm") or norm_text(c.get("consumer_name", ""))
        if qt and nm.startswith(qt):
            return (1, 0, nm)
        if qt and any(w.startswith(qt) for w in nm.split()):
            return (1, 1, nm)
        if qt and qt in nm:
            return (1, 2, nm)
        return (1, 3, nm)

    results.sort(key=_rank)
    return {"results": [consumer_view(c, user) for c in results], "exact_count": len(exact)}


@phed_router.get("/consumers")
async def list_consumers(
    page: int = 1, limit: int = 25, search: Optional[str] = None, ward_id: Optional[str] = None, colony: Optional[str] = None,
    service: Optional[str] = None, category: Optional[str] = None, status: Optional[str] = None, linked: Optional[str] = None,
    source: Optional[str] = None, user: dict = Depends(get_current_user),
):
    require_officer(user)
    db = get_db()
    q: Dict[str, Any] = {"is_active": True}
    if ward_id: q["ward_id"] = ward_id
    if colony: q["colony_name"] = {"$regex": f"^{re.escape(colony)}$", "$options": "i"}
    if category: q["category"] = category
    if status: q["status"] = status
    if source: q["source"] = source
    if linked == "yes": q["linked_property_id"] = {"$ne": None}
    if linked == "no": q["linked_property_id"] = None
    if service in SERVICES:
        refs = await db.phed_connections.distinct("consumer_ref", {"service": service, "is_active": True})
        q["id"] = {"$in": refs}
    if search:
        qn, qt = norm_id(search), norm_text(search)
        conn_refs = await db.phed_connections.distinct("consumer_ref", {"connection_number_norm": {"$regex": re.escape(qn)}})
        rx = {"$regex": re.escape(qt), "$options": "i"}
        q["$or"] = [{"consumer_id_norm": {"$regex": re.escape(qn)}}, {"phone_norm": {"$regex": re.escape(qn)}}, {"name_norm": rx}, {"address_norm": rx}, {"fh_name": rx}, {"locality": rx}, {"id": {"$in": conn_refs}}, {"linked_property_number": rx}]
    total = await db.phed_consumers.count_documents(q)
    items = await db.phed_consumers.find(q, {"_id": 0, "original_values": 0}).sort([("ward_number", 1), ("consumer_id", 1)]).skip((page - 1) * limit).limit(limit).to_list(None)
    await attach_connections(items)
    prop_ids = [c["linked_property_id"] for c in items if c.get("linked_property_id")]
    props = {p["id"]: p for p in await db.properties.find({"id": {"$in": prop_ids}}, {"_id": 0, "id": 1, "property_id": 1, "owner_name": 1, "assigned_employee_id": 1, "assigned_employee_ids": 1}).to_list(None)} if prop_ids else {}
    surveys = {s["property_record_id"]: s for s in await db.phed_surveys.find({"property_record_id": {"$in": prop_ids}, "status": {"$ne": "Rejected"}}, {"_id": 0, "property_record_id": 1, "status": 1, "surveyor_name": 1}).to_list(None)} if prop_ids else {}
    for c in items:
        p = props.get(c.get("linked_property_id"))
        c["linked_property"] = p
        s = surveys.get(c.get("linked_property_id"))
        c["survey_status"] = s["status"] if s else ("Imported" if c["source"] == "imported" else "Not Started")
        c["assigned_surveyor"] = s["surveyor_name"] if s else None
    return {"consumers": items, "total": total, "page": page, "pages": (total + limit - 1) // limit}


@phed_router.get("/consumers-by-locality")
async def consumers_by_locality(ward_id: Optional[str] = None, user: dict = Depends(get_current_user)):
    """Colony/locality-wise breakdown of consumers with counts (total, linked, unlinked)."""
    require_officer(user)
    db = get_db()
    match: Dict[str, Any] = {"is_active": True}
    if ward_id:
        match["ward_id"] = ward_id
    pipeline = [
        {"$match": match},
        {"$group": {
            "_id": {"$toUpper": {"$trim": {"input": {"$ifNull": ["$colony_name", ""]}}}},
            "count": {"$sum": 1},
            "linked": {"$sum": {"$cond": [{"$ifNull": ["$linked_property_id", False]}, 1, 0]}},
        }},
        {"$sort": {"count": -1}},
    ]
    rows = await db.phed_consumers.aggregate(pipeline).to_list(None)
    localities = [
        {"locality": (r["_id"] or "(No locality)"), "count": r["count"], "linked": r["linked"], "unlinked": r["count"] - r["linked"]}
        for r in rows
    ]
    return {"localities": localities, "total_localities": len(localities), "total_consumers": sum(r["count"] for r in rows)}


@phed_router.get("/consumers/{consumer_ref}")
async def get_consumer(consumer_ref: str, user: dict = Depends(get_current_user)):
    c = await get_db().phed_consumers.find_one({"id": consumer_ref}, {"_id": 0})
    if not c:
        raise HTTPException(404, "Consumer not found")
    await attach_connections([c])
    c["linked_property"] = await property_brief(c.get("linked_property_id"))
    c["surveys"] = await get_db().phed_surveys.find({"$or": [{"consumer_ref": consumer_ref}, {"consumer_refs": consumer_ref}]}, {"_id": 0, "attachments": 0}).sort("created_at", -1).to_list(20)
    if is_officer(user):
        c["audit"] = await get_db().phed_audit_logs.find({"entity_id": consumer_ref}, {"_id": 0}).sort("timestamp", -1).limit(50).to_list(None)
    return consumer_view(c, user)


@phed_router.post("/consumers")
async def create_consumer(data: ConsumerIn, user: dict = Depends(get_current_user)):
    if not data.consumer_name.strip():
        raise HTTPException(400, "Consumer name is required")
    db = get_db()
    if data.linked_property_id:
        await get_property_or_403(data.linked_property_id, user)
    cid = data.consumer_id.strip() if data.consumer_id else ""
    if cid and await db.phed_consumers.find_one({"consumer_id_norm": norm_id(cid)}):
        raise HTTPException(409, "A consumer with this Consumer ID already exists; search and link it instead")
    ward = await db.phed_wards.find_one({"id": data.ward_id}, {"_id": 0}) if data.ward_id else None
    cat, _ = norm_category(data.category or "")
    ts = now_iso()
    ref = str(uuid.uuid4())
    prov = None if cid else f"PROV{uuid.uuid4().hex[:8].upper()}"
    prop = await property_brief(data.linked_property_id)
    doc = {
        "id": ref, "consumer_id": cid or prov, "consumer_id_norm": norm_id(cid or prov), "provisional_ref": prov, "consumer_name": data.consumer_name.strip(),
        "name_norm": norm_text(data.consumer_name), "fh_name": (data.fh_name or "").strip(), "ppp_id": (data.ppp_id or "").strip(), "address": (data.address or "").strip(),
        "address_norm": norm_text(data.address or ""), "locality": (data.locality or "").strip(), "colony_name": (data.locality or (prop or {}).get("colony") or "").strip(),
        "phone": (data.phone or "").strip(), "phone_norm": norm_id(data.phone or ""), "ward_id": ward["id"] if ward else None, "ward_number": ward["ward_number"] if ward else (prop or {}).get("ward"),
        "category": cat, "category_raw": data.category or "", "source": "survey" if not is_admin(user) else "admin", "source_file": None, "source_row": None, "import_id": None,
        "linked_property_id": data.linked_property_id, "linked_property_number": (prop or {}).get("property_id"), "status": "Not Started", "is_active": True,
        "created_by": user["id"], "created_at": ts, "updated_at": ts,
    }
    await db.phed_consumers.insert_one(doc)
    doc.pop("_id", None)
    created = []
    for cn in data.connections or []:
        if cn.get("connection_number") and cn.get("service") in SERVICES:
            created.append(await _add_connection(ref, doc, ConnectionIn(**{**cn, "linked_property_id": data.linked_property_id}), user, "survey" if not is_admin(user) else "admin"))
    await audit(user, "CONSUMER_CREATE", "consumer", ref, after=doc)
    doc["connections"] = created
    return doc


@phed_router.put("/consumers/{consumer_ref}")
async def update_consumer(consumer_ref: str, data: ConsumerIn, user: dict = Depends(get_current_user)):
    db = get_db()
    before = await db.phed_consumers.find_one({"id": consumer_ref}, {"_id": 0})
    if not before:
        raise HTTPException(404, "Consumer not found")
    if not is_officer(user) and before.get("source") == "imported" and before.get("created_by") != user["id"]:
        # surveyors may only complete missing fields on imported master data
        allowed = {k: getattr(data, k) for k in ("fh_name", "ppp_id", "phone", "address", "locality") if getattr(data, k) and not before.get(k)}
        if not allowed:
            raise HTTPException(403, "Surveyors can only fill missing fields on imported records")
        upd = {**allowed, "updated_at": now_iso()}
    else:
        cat, _ = norm_category(data.category or before.get("category_raw") or "")
        upd = {"consumer_name": data.consumer_name.strip(), "name_norm": norm_text(data.consumer_name), "fh_name": (data.fh_name or "").strip(), "ppp_id": (data.ppp_id or "").strip(),
               "address": (data.address or "").strip(), "address_norm": norm_text(data.address or ""), "locality": (data.locality or "").strip(), "phone": (data.phone or "").strip(),
               "phone_norm": norm_id(data.phone or ""), "category": cat, "updated_at": now_iso()}
        if data.ward_id:
            ward = await db.phed_wards.find_one({"id": data.ward_id}, {"_id": 0})
            if ward:
                upd.update({"ward_id": ward["id"], "ward_number": ward["ward_number"]})
        if is_admin(user) and data.consumer_id and norm_id(data.consumer_id) != before["consumer_id_norm"]:
            if await db.phed_consumers.find_one({"consumer_id_norm": norm_id(data.consumer_id)}):
                raise HTTPException(409, "Consumer ID already in use")
            upd.update({"consumer_id": data.consumer_id.strip(), "consumer_id_norm": norm_id(data.consumer_id), "provisional_ref": None})
            await db.phed_connections.update_many({"consumer_ref": consumer_ref}, {"$set": {"consumer_id": data.consumer_id.strip()}})
    if "phone" in upd:
        upd["phone_norm"] = norm_id(upd["phone"])
    await db.phed_consumers.update_one({"id": consumer_ref}, {"$set": upd})
    await audit(user, "CONSUMER_UPDATE", "consumer", consumer_ref, before=before, after={**before, **upd})
    return {**before, **upd}


async def _add_connection(consumer_ref: str, consumer: dict, data: ConnectionIn, user: dict, source: str) -> dict:
    if data.service not in SERVICES:
        raise HTTPException(400, "Service must be Water or Sewer")
    cn = data.connection_number.strip()
    if not cn:
        raise HTTPException(400, "Connection number is required")
    db = get_db()
    dup = await db.phed_connections.find_one({"consumer_ref": consumer_ref, "service": data.service, "connection_number_norm": norm_id(cn), "is_active": True}, {"_id": 0})
    if dup:
        raise HTTPException(409, f"{data.service} connection {cn} already exists for this consumer")
    cat, _ = norm_category(data.category or consumer.get("category") or "")
    ts = now_iso()
    doc = {"id": str(uuid.uuid4()), "consumer_ref": consumer_ref, "consumer_id": consumer["consumer_id"], "connection_number": cn, "connection_number_norm": norm_id(cn),
           "service": data.service, "category": cat, "category_raw": data.category or "", "source": source, "status": "Active", "is_active": True,
           "verification_status": "Unverified", "verified_at": None, "verified_by": None, "verified_by_name": None, "remarks": (data.remarks or "").strip() or None,
           "linked_property_id": data.linked_property_id or consumer.get("linked_property_id"), "created_by": user["id"], "created_at": ts, "updated_at": ts}
    await db.phed_connections.insert_one(doc)
    doc.pop("_id", None)
    await audit(user, "CONNECTION_CREATE", "consumer", consumer_ref, after=doc)
    return doc


@phed_router.post("/consumers/{consumer_ref}/connections")
async def add_connection(consumer_ref: str, data: ConnectionIn, user: dict = Depends(get_current_user)):
    consumer = await get_db().phed_consumers.find_one({"id": consumer_ref}, {"_id": 0})
    if not consumer:
        raise HTTPException(404, "Consumer not found")
    if data.linked_property_id:
        await get_property_or_403(data.linked_property_id, user)
    return await _add_connection(consumer_ref, consumer, data, user, "survey" if not is_admin(user) else "admin")


@phed_router.delete("/connections/{connection_id}")
async def deactivate_connection(connection_id: str, user: dict = Depends(get_current_user)):
    require_admin(user)
    before = await get_db().phed_connections.find_one({"id": connection_id}, {"_id": 0})
    if not before:
        raise HTTPException(404, "Connection not found")
    await get_db().phed_connections.update_one({"id": connection_id}, {"$set": {"is_active": False, "status": "Removed", "updated_at": now_iso()}})
    await audit(user, "CONNECTION_REMOVE", "consumer", before["consumer_ref"], before=before)
    return {"message": "Connection removed"}


class DeleteAllIn(BaseModel):
    confirm: str
    ward_id: Optional[str] = None


@phed_router.post("/consumers/delete-all")
async def delete_all_consumers(data: DeleteAllIn, user: dict = Depends(get_current_user)):
    require_admin(user)
    if data.confirm != "DELETE":
        raise HTTPException(400, "Type DELETE to confirm")
    db = get_db()
    q = {"ward_id": data.ward_id} if data.ward_id else {}
    refs = [c["id"] for c in await db.phed_consumers.find(q, {"_id": 0, "id": 1}).to_list(None)]
    conn_res = await db.phed_connections.delete_many({"consumer_ref": {"$in": refs}} if data.ward_id else {})
    cons_res = await db.phed_consumers.delete_many(q)
    await clear_map_cache()
    await audit(user, "CONSUMERS_DELETE_ALL", "consumer", data.ward_id or "all",
                meta={"consumers": cons_res.deleted_count, "connections": conn_res.deleted_count})
    return {"message": "Consumers deleted", "deleted_consumers": cons_res.deleted_count, "deleted_connections": conn_res.deleted_count}


@phed_router.delete("/consumers/{consumer_ref}")
async def delete_consumer(consumer_ref: str, user: dict = Depends(get_current_user)):
    require_admin(user)
    db = get_db()
    before = await db.phed_consumers.find_one({"id": consumer_ref}, {"_id": 0})
    if not before:
        raise HTTPException(404, "Consumer not found")
    conn_res = await db.phed_connections.delete_many({"consumer_ref": consumer_ref})
    await db.phed_consumers.delete_one({"id": consumer_ref})
    await clear_map_cache()
    await audit(user, "CONSUMER_DELETE", "consumer", consumer_ref, before=before, meta={"connections": conn_res.deleted_count})
    return {"message": "Consumer deleted", "deleted_connections": conn_res.deleted_count}


class LocationIn(BaseModel):
    latitude: float = Field(..., ge=-90, le=90)
    longitude: float = Field(..., ge=-180, le=180)


@phed_router.put("/properties/{property_record_id}/location")
async def update_property_location(property_record_id: str, data: LocationIn, user: dict = Depends(get_current_user)):
    """Move a property pin on the map (surveyor: own/assigned properties; admin/officer: any)."""
    prop = await get_property_or_403(property_record_id, user)
    before = {"latitude": prop.get("latitude"), "longitude": prop.get("longitude")}
    after = {"latitude": data.latitude, "longitude": data.longitude}
    await get_db().properties.update_one({"id": property_record_id}, {"$set": {
        **after, "location_updated_by": user["id"], "location_updated_at": now_iso()}})
    await clear_map_cache()
    await audit(user, "PROPERTY_MOVE", "property", property_record_id, before=before, after=after)
    return {"message": "Location saved", **after}


@phed_router.post("/consumers/{consumer_ref}/link")
async def link_consumer(consumer_ref: str, data: LinkIn, user: dict = Depends(get_current_user)):
    if not data.confirm:
        raise HTTPException(400, "Explicit confirmation is required to link a consumer to a property")
    db = get_db()
    consumer = await db.phed_consumers.find_one({"id": consumer_ref}, {"_id": 0})
    if not consumer:
        raise HTTPException(404, "Consumer not found")
    prop = await get_property_or_403(data.property_record_id, user)
    if consumer.get("linked_property_id") and consumer["linked_property_id"] != prop["id"] and not is_admin(user):
        raise HTTPException(409, "Consumer is already linked to another property; ask an administrator to correct it")
    upd = {"linked_property_id": prop["id"], "linked_property_number": prop["property_id"], "updated_at": now_iso()}
    if not consumer.get("colony_name"):
        upd["colony_name"] = prop.get("colony")
    await db.phed_consumers.update_one({"id": consumer_ref}, {"$set": upd})
    await db.phed_connections.update_many({"consumer_ref": consumer_ref, "linked_property_id": None}, {"$set": {"linked_property_id": prop["id"]}})
    await audit(user, "CONSUMER_LINK", "consumer", consumer_ref, before={"linked_property_id": consumer.get("linked_property_id")}, after=upd, meta={"property_id": prop["property_id"]})
    return {**consumer, **upd, "linked_property": await property_brief(prop["id"])}


@phed_router.post("/consumers/{consumer_ref}/unlink")
async def unlink_consumer(consumer_ref: str, user: dict = Depends(get_current_user)):
    require_admin(user)
    db = get_db()
    consumer = await db.phed_consumers.find_one({"id": consumer_ref}, {"_id": 0})
    if not consumer:
        raise HTTPException(404, "Consumer not found")
    await db.phed_consumers.update_one({"id": consumer_ref}, {"$set": {"linked_property_id": None, "linked_property_number": None, "updated_at": now_iso()}})
    await db.phed_connections.update_many({"consumer_ref": consumer_ref}, {"$set": {"linked_property_id": None}})
    await audit(user, "CONSUMER_UNLINK", "consumer", consumer_ref, before={"linked_property_id": consumer.get("linked_property_id"), "linked_property_number": consumer.get("linked_property_number")})
    return {"message": "Consumer unlinked from property"}


# ---------------- bulk consumer -> property auto-link ----------------
_STOP_TOKENS = {"h", "no", "hno", "house", "near", "opp", "gali", "st", "road", "rd", "ward", "the", "and", "colony", "nagar"}


def _tokens(s: str):
    return {t for t in norm_text(s).replace(",", " ").replace(".", " ").replace("-", " ").split() if len(t) >= 2 and t not in _STOP_TOKENS}


def _match_score(c_addr, c_name, p_addr, p_name):
    ca, pa = _tokens(c_addr), _tokens(p_addr)
    cn, pn = _tokens(c_name), _tokens(p_name)
    addr_overlap = len(ca & pa)
    name_overlap = len(cn & pn)
    if name_overlap >= 1 and addr_overlap >= 2:
        return addr_overlap * 2 + name_overlap * 3, "high"
    if addr_overlap >= 2 or (name_overlap >= 1 and addr_overlap >= 1):
        return addr_overlap * 2 + name_overlap, "medium"
    return 0, None


@phed_router.get("/link-suggestions")
async def link_suggestions(limit: int = 500, user: dict = Depends(get_current_user)):
    require_admin(user)
    db = get_db()
    consumers = await db.phed_consumers.find(
        {"is_active": True, "$or": [{"linked_property_id": None}, {"linked_property_id": {"$exists": False}}]},
        {"_id": 0, "id": 1, "consumer_id": 1, "consumer_name": 1, "fh_name": 1, "address": 1, "locality": 1, "phone": 1, "phone_norm": 1}).limit(50000).to_list(None)
    total_unlinked = len(consumers)
    props = await db.properties.find({}, {"_id": 0, "id": 1, "property_id": 1, "owner_name": 1, "address": 1, "colony": 1, "mobile": 1}).to_list(None)
    # phone index (normalized property mobile -> property index) + address token inverted index
    phone_idx: Dict[str, int] = {}
    inv: Dict[str, set] = {}
    for i, p in enumerate(props):
        pm = norm_id(p.get("mobile", ""))
        if len(pm) >= 7 and len(set(pm)) > 1:  # skip junk like 0000000000 / 9999999999
            phone_idx.setdefault(pm, i)
        for t in _tokens(p.get("address", "")) | _tokens(p.get("colony", "")):
            inv.setdefault(t, set()).add(i)
    suggestions = []
    for c in consumers:
        best = None
        # 1) strongest signal: consumer phone == property mobile
        cph = c.get("phone_norm") or norm_id(c.get("phone", ""))
        if len(cph) >= 7 and len(set(cph)) > 1 and cph in phone_idx:
            best = {"score": 100, "confidence": "high", "reason": "phone match", "property": props[phone_idx[cph]]}
        # 2) fallback: address + owner-name token overlap
        if best is None:
            c_addr = f"{c.get('address', '')} {c.get('locality', '')}"
            cand_idx = set()
            for t in _tokens(c_addr) | _tokens(c.get("consumer_name", "")):
                cand_idx |= inv.get(t, set())
            for i in cand_idx:
                p = props[i]
                score, conf = _match_score(c_addr, c.get("consumer_name", ""), f"{p.get('address', '')} {p.get('colony', '')}", p.get("owner_name", ""))
                if conf and (best is None or score > best["score"]):
                    best = {"score": score, "confidence": conf, "reason": "address & name", "property": p}
        if best:
            suggestions.append({
                "consumer_ref": c["id"], "consumer_id": c.get("consumer_id"), "consumer_name": c.get("consumer_name"),
                "consumer_address": c.get("address"), "consumer_locality": c.get("locality"), "consumer_phone": c.get("phone"),
                "property_record_id": best["property"]["id"], "property_id": best["property"].get("property_id"),
                "property_owner": best["property"].get("owner_name"), "property_address": best["property"].get("address"),
                "confidence": best["confidence"], "score": best["score"], "reason": best["reason"],
            })
    suggestions.sort(key=lambda x: (0 if x["confidence"] == "high" else 1, -x["score"]))
    return {"total_unlinked": total_unlinked, "properties_available": len(props), "suggestions": suggestions[:limit]}


class BulkLinkItem(BaseModel):
    consumer_ref: str
    property_record_id: str


class BulkLinkIn(BaseModel):
    links: List[BulkLinkItem]


@phed_router.post("/link-suggestions/apply")
async def apply_link_suggestions(data: BulkLinkIn, user: dict = Depends(get_current_user)):
    require_admin(user)
    db = get_db()
    linked, skipped, errors = 0, 0, []
    for item in data.links:
        consumer = await db.phed_consumers.find_one({"id": item.consumer_ref}, {"_id": 0})
        prop = await db.properties.find_one({"id": item.property_record_id}, {"_id": 0})
        if not consumer or not prop:
            skipped += 1; errors.append({"consumer_ref": item.consumer_ref, "reason": "consumer or property not found"}); continue
        upd = {"linked_property_id": prop["id"], "linked_property_number": prop["property_id"], "updated_at": now_iso()}
        if not consumer.get("colony_name"):
            upd["colony_name"] = prop.get("colony")
        await db.phed_consumers.update_one({"id": item.consumer_ref}, {"$set": upd})
        await db.phed_connections.update_many({"consumer_ref": item.consumer_ref, "linked_property_id": None}, {"$set": {"linked_property_id": prop["id"]}})
        await audit(user, "CONSUMER_LINK", "consumer", item.consumer_ref, before={"linked_property_id": consumer.get("linked_property_id")}, after=upd, meta={"property_id": prop["property_id"], "via": "bulk_auto_link"})
        linked += 1
    return {"linked": linked, "skipped": skipped, "errors": errors}



# ---------------- property PHED section ----------------
@phed_router.get("/property/{property_record_id}")
async def property_phed_section(property_record_id: str, user: dict = Depends(get_current_user)):
    prop = await get_property_or_403(property_record_id, user)
    db = get_db()
    consumers = await db.phed_consumers.find({"linked_property_id": property_record_id, "is_active": True}, {"_id": 0, "original_values": 0}).to_list(None)
    await attach_connections(consumers)
    survey = await db.phed_surveys.find_one({"property_record_id": property_record_id, "status": {"$ne": "Rejected"}}, {"_id": 0})
    legacy = await db.submissions.find_one({"property_record_id": property_record_id, "status": {"$ne": "Rejected"}}, {"_id": 0, "id": 1, "status": 1, "submitted_at": 1, "receiver_name": 1})
    return {"property": prop, "consumers": [consumer_view(c, user) for c in consumers], "survey": survey, "legacy_submission": legacy}


# ---------------- surveys ----------------
class SurveyDraftIn(BaseModel):
    property_record_id: str
    survey_type: Optional[str] = None
    consumer_refs: Optional[List[str]] = None
    latitude: Optional[float] = None
    longitude: Optional[float] = None
    gps_accuracy: Optional[float] = None
    gps_captured_at: Optional[str] = None
    polygon: Optional[List[List[float]]] = None       # building outline [[lat,lng], ...]
    surveyor_latitude: Optional[float] = None          # raw GPS where surveyor stood
    surveyor_longitude: Optional[float] = None
    remarks: Optional[str] = None
    corrections: Optional[Dict[str, str]] = None  # surveyor-noted corrections (not applied to MC master)
    owner_mismatch: Optional[bool] = None            # property owner != PHED bill owner (neither record is changed)
    mismatch_reason: Optional[str] = None
    verified_connection_ids: Optional[List[str]] = None  # connections physically verified on site
    connection_remarks: Optional[Dict[str, str]] = None  # {connection_id: remark}
    water: Optional[Dict[str, Any]] = None  # quick water-supply survey answers (surveyor simple form)


def survey_public(s: dict) -> dict:
    s = dict(s)
    for a in s.get("attachments", []):
        a.pop("file_id", None)
    return s


@phed_router.post("/surveys/draft")
async def save_draft(data: SurveyDraftIn, user: dict = Depends(get_current_user)):
    prop = await get_property_or_403(data.property_record_id, user)
    db = get_db()
    if data.survey_type and data.survey_type not in SURVEY_TYPES:
        raise HTTPException(400, "Invalid survey type")
    existing = await db.phed_surveys.find_one({"property_record_id": prop["id"], "status": {"$ne": "Rejected"}}, {"_id": 0})
    if existing and existing["status"] == "Approved":
        raise HTTPException(409, "Survey already approved for this property")
    ts = now_iso()
    poly = data.polygon or None
    centroid = polygon_centroid(poly) if poly and len(poly) >= 3 else None
    # House-center location: polygon centroid takes priority; else explicit lat/long.
    if centroid:
        house_lat, house_lng = centroid[0], centroid[1]
    else:
        house_lat, house_lng = data.latitude, data.longitude
    sv_lat, sv_lng = data.surveyor_latitude, data.surveyor_longitude
    # Informational distance surveyor stood from house center (NOT flagged as error).
    surveyor_dist = None
    if sv_lat is not None and sv_lng is not None and house_lat is not None and house_lng is not None:
        surveyor_dist = round(haversine_m(sv_lat, sv_lng, house_lat, house_lng), 1)
    # Legacy drift vs imported MC point; only flag when NO polygon was drawn.
    drift = None
    ref_lat = prop.get("mc_latitude") if prop.get("mc_latitude") is not None else prop.get("latitude")
    ref_lng = prop.get("mc_longitude") if prop.get("mc_longitude") is not None else prop.get("longitude")
    if house_lat is not None and house_lng is not None and ref_lat and ref_lng:
        drift = round(haversine_m(house_lat, house_lng, float(ref_lat), float(ref_lng)), 1)
    flagged = bool(not poly and drift and drift > GPS_DRIFT_FLAG_METERS)
    # Persist the drawn polygon + centroid onto the property so it can be reused/corrected.
    # The property's original MC latitude/longitude are NEVER overwritten.
    if poly and centroid:
        prop_upd = {"polygon": poly, "polygon_centroid": {"latitude": centroid[0], "longitude": centroid[1]},
                    "polygon_source": prop.get("polygon_source") or "surveyor", "polygon_updated_at": ts}
        await db.properties.update_one({"id": prop["id"]}, {"$set": prop_upd})
        prop = {**prop, **prop_upd}
        await clear_map_cache()
    ward = await db.phed_wards.find_one({"colonies": prop.get("colony")}, {"_id": 0}) if prop.get("colony") else None
    fields = {
        "survey_type": data.survey_type, "consumer_refs": data.consumer_refs or [], "consumer_ref": (data.consumer_refs or [None])[0],
        "latitude": house_lat, "longitude": house_lng, "gps_accuracy": data.gps_accuracy, "gps_captured_at": data.gps_captured_at,
        "polygon": poly, "polygon_source": ("surveyor" if poly else None),
        "surveyor_latitude": sv_lat, "surveyor_longitude": sv_lng, "surveyor_distance_m": surveyor_dist,
        "property_latitude": prop.get("latitude"), "property_longitude": prop.get("longitude"), "gps_drift_m": drift, "gps_flagged": flagged,
        "remarks": data.remarks, "corrections": data.corrections or {}, "updated_at": ts,
        "owner_mismatch": bool(data.owner_mismatch), "mismatch_reason": (data.mismatch_reason or "").strip() or None,
        "verified_connection_ids": data.verified_connection_ids, "connection_remarks": data.connection_remarks or {},
        "water": data.water or {},
    }
    if existing:
        await db.phed_surveys.update_one({"id": existing["id"]}, {"$set": fields})
        return survey_public({**existing, **fields})
    doc = {"id": str(uuid.uuid4()), "property_record_id": prop["id"], "property_id": prop["property_id"], "surveyor_id": user["id"], "surveyor_name": user["name"],
           "ward_id": ward["id"] if ward else None, "ward_number": ward["ward_number"] if ward else prop.get("ward"), "colony_name": prop.get("colony"),
           "status": "Draft", "attachments": [], "started_at": ts, "submitted_at": None, "rejection_reason": None, "created_at": ts, **fields}
    await db.phed_surveys.insert_one(doc)
    doc.pop("_id", None)
    # PHED status lives in its own field; the legacy property-survey status is untouched.
    await db.properties.update_one({"id": prop["id"]}, {"$set": {"phed_survey_status": "Draft"}})
    await clear_map_cache()
    return survey_public(doc)


@phed_router.post("/surveys/{survey_id}/attachments")
async def upload_attachment(survey_id: str, attachment_type: str = Form(...), file: UploadFile = File(...),
                            latitude: Optional[float] = Form(None), longitude: Optional[float] = Form(None),
                            captured_at: Optional[str] = Form(None), gps_source: Optional[str] = Form(None),
                            user: dict = Depends(get_current_user)):
    db = get_db()
    survey = await db.phed_surveys.find_one({"id": survey_id}, {"_id": 0})
    if not survey:
        raise HTTPException(404, "Survey not found")
    if survey["surveyor_id"] != user["id"] and not is_admin(user):
        raise HTTPException(403, "Not your survey")
    if survey["status"] == "Approved":
        raise HTTPException(409, "Survey already approved")
    if attachment_type not in ATTACHMENT_TYPES:
        raise HTTPException(400, "Invalid attachment type")
    content = await file.read()
    if len(content) > MAX_ATTACHMENT_BYTES:
        raise HTTPException(413, "File too large (max 20 MB)")
    ctype = (file.content_type or "").lower()
    if ctype not in ALLOWED_MIME:
        raise HTTPException(400, "Only JPEG/PNG/WebP images or PDF are allowed")
    ext = {"image/jpeg": "jpg", "image/png": "png", "image/webp": "webp", "application/pdf": "pdf"}[ctype]
    safe_name = f"phed_{survey['property_id']}_{attachment_type.lower()}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.{ext}"
    file_id = await save_file_to_gridfs(content, safe_name, ctype)
    att = {"id": str(uuid.uuid4()), "attachment_type": attachment_type, "file_id": file_id, "filename": safe_name, "content_type": ctype, "size": len(content),
           "uploaded_by": user["id"], "uploaded_at": now_iso(),
           "latitude": latitude, "longitude": longitude, "captured_at": captured_at, "gps_source": gps_source}
    await db.phed_surveys.update_one({"id": survey_id}, {"$push": {"attachments": att}, "$set": {"updated_at": now_iso()}})
    await audit(user, "DOCUMENT_UPLOAD", "survey", survey_id, meta={"attachment_type": attachment_type, "attachment_id": att["id"]})
    return {k: v for k, v in att.items() if k != "file_id"}


@phed_router.delete("/surveys/{survey_id}/attachments/{attachment_id}")
async def delete_attachment(survey_id: str, attachment_id: str, user: dict = Depends(get_current_user)):
    db = get_db()
    survey = await db.phed_surveys.find_one({"id": survey_id}, {"_id": 0})
    if not survey:
        raise HTTPException(404, "Survey not found")
    if survey["surveyor_id"] != user["id"] and not is_admin(user):
        raise HTTPException(403, "Not your survey")
    if survey["status"] == "Approved" and not is_admin(user):
        raise HTTPException(409, "Survey already approved")
    att = next((a for a in survey["attachments"] if a["id"] == attachment_id), None)
    if not att:
        raise HTTPException(404, "Attachment not found")
    try:
        await get_fs().delete(ObjectId(att["file_id"]))
    except Exception:
        pass
    await db.phed_surveys.update_one({"id": survey_id}, {"$pull": {"attachments": {"id": attachment_id}}})
    await audit(user, "DOCUMENT_REMOVE", "survey", survey_id, meta={"attachment_id": attachment_id})
    return {"message": "Attachment removed"}


@phed_router.get("/surveys/{survey_id}/attachments/{attachment_id}")
async def get_attachment(survey_id: str, attachment_id: str, user: dict = Depends(get_current_user)):
    survey = await get_db().phed_surveys.find_one({"id": survey_id}, {"_id": 0})
    if not survey:
        raise HTTPException(404, "Survey not found")
    if survey["surveyor_id"] != user["id"] and not is_officer(user):
        raise HTTPException(403, "Access denied")
    att = next((a for a in survey["attachments"] if a["id"] == attachment_id), None)
    if not att:
        raise HTTPException(404, "Attachment not found")
    grid_out = await get_fs().open_download_stream(ObjectId(att["file_id"]))
    content = await grid_out.read()
    return Response(content, media_type=att["content_type"], headers={"Cache-Control": "private, no-store", "Content-Disposition": f'inline; filename="{att["filename"]}"'})


# Print-ready label + ordering for the combined PDF
_DOC_ORDER = ["APPLICATION", "AADHAAR_FRONT", "AADHAAR_BACK", "AADHAAR", "PROPERTY_PROOF",
              "PROPERTY_FRONT", "HOUSE_PHOTO", "DEATH_CERTIFICATE", "REGISTRY", "BILL", "OTHER"]
_DOC_TITLE = {
    "APPLICATION": "Application (आवेदन)", "AADHAAR": "Aadhaar Card", "AADHAAR_FRONT": "Aadhaar Card",
    "AADHAAR_BACK": "Aadhaar Card", "PROPERTY_PROOF": "Property Proof / Registry", "PROPERTY_FRONT": "Property Photo",
    "HOUSE_PHOTO": "House Photo (with owner)", "DEATH_CERTIFICATE": "Death Certificate", "REGISTRY": "Registry",
    "BILL": "Water / Sewer Bill", "OTHER": "Other Document",
}


async def _build_documents_pdf(survey: dict):
    """Combined print-ready HD PDF (bytes) or None if the survey has no documents.
    Aadhaar front+back share one page; Property proof spans consecutive pages; others one page each."""
    import fitz  # PyMuPDF
    atts = survey.get("attachments", []) or []
    if not atts:
        return None

    async def load(att):
        g = await get_fs().open_download_stream(ObjectId(att["file_id"]))
        return await g.read()

    by_type: Dict[str, list] = {}
    for a in atts:
        by_type.setdefault(a["attachment_type"], []).append(a)
    for k in by_type:
        by_type[k].sort(key=lambda a: a.get("uploaded_at", ""))

    out = fitz.open()
    PW, PH, M = 595.0, 842.0, 30.0

    def header(page, text):
        page.insert_text((M, 22), text, fontsize=10, color=(0.2, 0.2, 0.2))

    def img_rect(top=False, bottom=False):
        if top:
            return fitz.Rect(M, 34, PW - M, M + (PH - 60) / 2 - 6)
        if bottom:
            return fitz.Rect(M, M + (PH - 60) / 2 + 6, PW - M, PH - M)
        return fitz.Rect(M, 34, PW - M, PH - M)

    async def add_doc(att):
        data = await load(att)
        if (att.get("content_type") or "").lower() == "application/pdf":
            src = fitz.open(stream=data, filetype="pdf")
            out.insert_pdf(src)
            src.close()
            return None
        return data

    ref = survey.get("reference_number") or ""
    prop_label = f"{ref + '  ·  ' if ref else ''}{survey.get('property_id', '')} · {survey.get('owner_name', '')}"
    ordered = [t for t in _DOC_ORDER if t in by_type] + [t for t in by_type if t not in _DOC_ORDER]

    for t in ordered:
        items = by_type[t]
        if t == "AADHAAR_FRONT":
            backs = by_type.get("AADHAAR_BACK", [])
            for i, fatt in enumerate(items):
                fbytes = await add_doc(fatt)
                page = out.new_page(width=PW, height=PH)
                header(page, f"Aadhaar Card — {prop_label}")
                if fbytes:
                    page.insert_image(img_rect(top=True), stream=fbytes, keep_proportion=True)
                    page.insert_text((M, 34 + (PH - 60) / 2 - 14), "Front", fontsize=8, color=(0.4, 0.4, 0.4))
                b = backs[i] if i < len(backs) else None
                if b:
                    bbytes = await add_doc(b)
                    if bbytes:
                        page.insert_image(img_rect(bottom=True), stream=bbytes, keep_proportion=True)
                        page.insert_text((M, M + (PH - 60) / 2 + 2), "Back", fontsize=8, color=(0.4, 0.4, 0.4))
        elif t == "AADHAAR_BACK":
            fronts = by_type.get("AADHAAR_FRONT", [])
            for b in items[len(fronts):]:
                bbytes = await add_doc(b)
                if bbytes:
                    page = out.new_page(width=PW, height=PH)
                    header(page, f"Aadhaar Card (Back) — {prop_label}")
                    page.insert_image(img_rect(), stream=bbytes, keep_proportion=True)
        else:
            title = _DOC_TITLE.get(t, t)
            for idx, att in enumerate(items):
                data = await add_doc(att)
                if data is None:
                    continue
                page = out.new_page(width=PW, height=PH)
                suffix = f" (page {idx + 1}/{len(items)})" if len(items) > 1 else ""
                header(page, f"{title}{suffix} — {prop_label}")
                page.insert_image(img_rect(), stream=data, keep_proportion=True)

    buf = out.tobytes(deflate=True)
    out.close()
    return buf


@phed_router.get("/surveys/{survey_id}/documents.pdf")
async def download_survey_documents(survey_id: str, user: dict = Depends(get_current_user)):
    """Combined, print-ready HD PDF of a survey's documents."""
    survey = await get_db().phed_surveys.find_one({"id": survey_id}, {"_id": 0})
    if not survey:
        raise HTTPException(404, "Survey not found")
    if survey["surveyor_id"] != user["id"] and not is_officer(user):
        raise HTTPException(403, "Access denied")
    buf = await _build_documents_pdf(survey)
    if buf is None:
        raise HTTPException(404, "No documents to download")
    fname = f"phed_documents_{survey.get('reference_number') or survey.get('property_id') or survey_id}.pdf"
    return Response(buf, media_type="application/pdf",
                    headers={"Cache-Control": "private, no-store", "Content-Disposition": f'attachment; filename="{fname}"'})


class BulkDocsIn(BaseModel):
    survey_ids: List[str] = []
    status: Optional[str] = None


@phed_router.post("/surveys/documents/bulk.zip")
async def bulk_download_documents(body: BulkDocsIn, user: dict = Depends(get_current_user)):
    """ZIP of per-survey combined PDFs. Pass survey_ids, or a status (defaults to Approved)."""
    require_officer(user)
    import zipfile
    db = get_db()
    q: Dict[str, Any] = {}
    if body.survey_ids:
        q["id"] = {"$in": body.survey_ids}
    if body.status:
        q["status"] = body.status
    if not body.survey_ids and not body.status:
        q["status"] = "Approved"
    surveys = await db.phed_surveys.find(q, {"_id": 0}).sort("updated_at", -1).to_list(500)
    if not surveys:
        raise HTTPException(404, "No surveys found")
    mem = io.BytesIO()
    count = 0
    used = set()
    with zipfile.ZipFile(mem, "w", zipfile.ZIP_DEFLATED) as zf:
        for s in surveys:
            pdf = await _build_documents_pdf(s)
            if not pdf:
                continue
            base = s.get("reference_number") or s.get("property_id") or s.get("id")
            name = f"{base}.pdf"
            n = 1
            while name in used:
                n += 1
                name = f"{base}_{n}.pdf"
            used.add(name)
            zf.writestr(name, pdf)
            count += 1
    if count == 0:
        raise HTTPException(404, "Selected surveys have no documents")
    mem.seek(0)
    return Response(mem.getvalue(), media_type="application/zip",
                    headers={"Cache-Control": "private, no-store", "Content-Disposition": f'attachment; filename="phed_documents_bulk_{count}.zip"'})


async def _next_reference(db, code: str) -> str:
    """Atomic, per-outcome-code sequence, e.g. DPS0001, SS0012."""
    doc = await db.phed_counters.find_one_and_update(
        {"code": code}, {"$inc": {"seq": 1}}, upsert=True, return_document=ReturnDocument.AFTER)
    return f"{code}{doc['seq']:04d}"


def _outcome_code(survey: dict, doc_pending: bool) -> str:
    """Reference-number prefix based on the survey outcome."""
    if doc_pending:
        return "DPS"                      # Document Pending Submission
    w = survey.get("water") or {}
    if w.get("property_locked"):
        return "PL"                       # Property Locked
    if w.get("owner_denied"):
        return "OD"                       # Owner Denied (assumption — not specified)
    oc = w.get("owner_change")
    if oc == "DEATH_TRANSFER":
        return "DS"                       # Death transfer Submission
    if oc == "OWNERSHIP_CHANGE":
        return "OC"                       # Ownership Change
    if survey.get("survey_type") == "NO_CONNECTION" or w.get("new_connection"):
        return "FS"                       # Final (new connection) Submission
    if w.get("has_connection"):
        has_sewer = bool(w.get("has_sewer") or (w.get("sewer_connection_numbers") or []))
        has_water = bool(w.get("connection_numbers") or [])
        if has_water and has_sewer:
            return "AS"                   # Already have both (direct submit)
        return "SS"                       # Sewer application Submission
    return "GS"                           # Generic Submission (fallback)


def _survey_outcome(survey: dict) -> str:
    """Survey outcome bucket for the PHED dashboard/map stats (not house-tax)."""
    w = survey.get("water") or {}
    if w.get("property_locked"):
        return "LOCKED"          # Property बंद / locked मिली
    if w.get("owner_denied"):
        return "DENIED"          # Owner ने मना किया
    if survey.get("survey_type") == "NO_CONNECTION" or w.get("new_connection"):
        return "NEW"             # New connection चाहिए
    if w.get("has_connection"):
        return "HAS_CONNECTION"  # पहले से connection है
    return "OTHER"


@phed_router.post("/surveys/{survey_id}/submit")
async def submit_survey(survey_id: str, user: dict = Depends(get_current_user)):
    db = get_db()
    survey = await db.phed_surveys.find_one({"id": survey_id}, {"_id": 0})
    if not survey:
        raise HTTPException(404, "Survey not found")
    if survey["surveyor_id"] != user["id"] and not is_admin(user):
        raise HTTPException(403, "Not your survey")
    if survey["status"] == "Approved":
        raise HTTPException(409, "Survey already approved")
    if survey.get("survey_type") not in SURVEY_TYPES:
        raise HTTPException(400, "Select the PHED connection status before submitting")
    if survey.get("latitude") is None or survey.get("longitude") is None:
        raise HTTPException(400, "GPS location is required before submitting")
    if (survey.get("water") or {}).get("document_pending"):
        missing = (survey.get("water") or {}).get("missing_documents") or []
        raise HTTPException(400, f"Required documents missing: {', '.join(missing) or 'documents'}")
    consumer_refs = survey.get("consumer_refs") or []
    if survey["survey_type"] in ("EXISTING_LINKED", "NEW_UNLISTED"):
        linked = await db.phed_consumers.count_documents({"linked_property_id": survey["property_record_id"], "is_active": True})
        if not consumer_refs and linked == 0:
            raise HTTPException(400, "Link or create at least one PHED consumer for this property")
        if not consumer_refs:
            consumer_refs = [c["id"] for c in await db.phed_consumers.find({"linked_property_id": survey["property_record_id"], "is_active": True}, {"_id": 0, "id": 1}).to_list(None)]
    ts = now_iso()
    if survey["survey_type"] == "WATER_CONNECTION":
        # Quick water-supply survey: record the bill details as a survey-sourced consumer + water connection(s)
        w = survey.get("water") or {}
        conn_numbers = [str(x).strip() for x in (w.get("connection_numbers") or ([w.get("connection_number")] if w.get("connection_number") else [])) if str(x).strip()]
        prop = await db.properties.find_one({"id": survey["property_record_id"]}, {"_id": 0, "owner_name": 1, "address": 1, "colony": 1, "mobile": 1})
        existing_c = await db.phed_consumers.find_one({"linked_property_id": survey["property_record_id"], "source": "survey", "is_active": True}, {"_id": 0})
        cdoc = {
            "consumer_name": (w.get("consumer_name") or (prop or {}).get("owner_name") or "").strip(),
            "fh_name": (w.get("fh_name") or "").strip(), "phone": str(w.get("phone") or (prop or {}).get("mobile") or "").strip(),
            "address": (prop or {}).get("address") or "", "locality": (prop or {}).get("colony") or "", "category": w.get("category") or "Domestic",
            "consumer_id": str(w.get("consumer_id") or "").strip(), "consumer_id_norm": norm_id(str(w.get("consumer_id") or "")),
            "ward_id": survey.get("ward_id"), "ward_number": survey.get("ward_number"), "colony_name": survey.get("colony_name"),
            "linked_property_id": survey["property_record_id"], "linked_property_number": survey["property_id"], "status": "Surveyed", "updated_at": ts,
        }
        if existing_c:
            cid = existing_c["id"]
            await db.phed_consumers.update_one({"id": cid}, {"$set": cdoc})
        else:
            cid = str(uuid.uuid4())
            prov = f"PROV-{survey['property_id']}"
            await db.phed_consumers.insert_one({"id": cid, "provisional_ref": prov, "source": "survey", "is_active": True, "created_by": user["id"], "created_at": ts,
                                                **({"consumer_id": prov, "consumer_id_norm": norm_id(prov)} if not cdoc["consumer_id"] else {}),
                                                **{k: v for k, v in cdoc.items() if k not in ("consumer_id", "consumer_id_norm") or cdoc["consumer_id"]}})
        for cn in conn_numbers:
            if not await db.phed_connections.find_one({"consumer_ref": cid, "connection_number_norm": norm_id(cn), "service": "Water", "is_active": True}):
                await db.phed_connections.insert_one({"id": str(uuid.uuid4()), "consumer_ref": cid, "consumer_id": cdoc["consumer_id"] or f"PROV-{survey['property_id']}",
                                                      "connection_number": cn, "connection_number_norm": norm_id(cn), "service": "Water", "category": cdoc["category"], "category_raw": cdoc["category"],
                                                      "source": "survey", "status": "Active", "is_active": True, "verification_status": "Verified", "verified_at": ts, "verified_by": user["id"],
                                                      "verified_by_name": user.get("name"), "remarks": None, "linked_property_id": survey["property_record_id"], "created_by": user["id"], "created_at": ts, "updated_at": ts})
        consumer_refs = [cid]
    # Submitted → admin approves. Anything needing a second look is parked in "Requires Review".
    w_meta = survey.get("water") or {}
    doc_pending = bool(w_meta.get("document_pending"))
    needs_review = bool(survey.get("gps_flagged") or survey.get("owner_mismatch") or survey["survey_type"] == "NEW_UNLISTED")
    if doc_pending:
        new_status = "Document Pending"
    else:
        new_status = "Requires Review" if needs_review else "Submitted"
    review_reasons = [r for r, on in (("GPS drift", survey.get("gps_flagged")), ("Owner mismatch", survey.get("owner_mismatch")),
                                      ("New/unlisted connection", survey["survey_type"] == "NEW_UNLISTED"),
                                      ("Documents pending", doc_pending)) if on]
    # Reference number — prefix reflects the outcome; regenerate when the outcome changes on re-submit
    code = _outcome_code(survey, doc_pending)
    ref = survey.get("reference_number")
    if not ref or survey.get("reference_code") != code:
        ref = await _next_reference(db, code)
    upd = {"status": new_status, "submitted_at": ts, "consumer_refs": consumer_refs, "consumer_ref": consumer_refs[0] if consumer_refs else None,
           "review_reasons": review_reasons, "reference_number": ref, "reference_code": code, "updated_at": ts}
    res = await db.phed_surveys.update_one({"id": survey_id, "status": {"$ne": "Approved"}}, {"$set": upd})
    if res.modified_count == 0:
        raise HTTPException(409, "Survey already approved")
    if consumer_refs:
        await db.phed_consumers.update_many({"id": {"$in": consumer_refs}}, {"$set": {"status": "Surveyed", "updated_at": ts}})
        conn_q: Dict[str, Any] = {"consumer_ref": {"$in": consumer_refs}, "is_active": True}
        verified_ids = survey.get("verified_connection_ids")
        if verified_ids is not None:
            conn_q["id"] = {"$in": verified_ids}
        await db.phed_connections.update_many(conn_q, {"$set": {"verification_status": "Verified", "verified_at": ts, "verified_by": user["id"],
                                                                "verified_by_name": user.get("name"), "linked_property_id": survey["property_record_id"], "updated_at": ts}})
        for cid, rem in (survey.get("connection_remarks") or {}).items():
            if rem:
                await db.phed_connections.update_one({"id": cid}, {"$set": {"remarks": rem}})
    phed_status = new_status if doc_pending else ("No PHED Connection" if survey["survey_type"] == "NO_CONNECTION" else new_status)
    await db.properties.update_one({"id": survey["property_record_id"]}, {"$set": {"phed_survey_status": phed_status, "phed_survey_type": survey["survey_type"], "phed_surveyed_at": ts,
                                                                                   "phed_survey_state": new_status, "phed_outcome": _survey_outcome(survey),
                                                                                   "phed_surveyor_id": survey.get("surveyor_id"), "phed_surveyor_name": survey.get("surveyor_name")}})
    await clear_map_cache()
    await audit(user, "SURVEY_SUBMIT", "survey", survey_id, after=upd, meta={"property_id": survey["property_id"], "survey_type": survey["survey_type"]})
    return {"message": "PHED survey submitted", "status": new_status, "survey_id": survey_id, "reference_number": ref, "reference_code": code}


class FieldPropertyIn(BaseModel):
    owner_name: str
    mobile: str = ""
    alternate_mobile: str = ""
    ward: str = ""
    address: str = ""
    colony: str = ""
    category: str = "Residential"
    latitude: float
    longitude: float


@phed_router.post("/field-properties")
async def create_field_property(body: FieldPropertyIn, user: dict = Depends(get_current_user)):
    """Surveyor creates a brand-new property point from the field (no MC property ID yet)."""
    db = get_db()
    if not body.owner_name.strip():
        raise HTTPException(400, "Owner name is required")
    ts = now_iso()
    pid = f"FIELD-{str(uuid.uuid4())[:8].upper()}"
    _colony = body.colony.strip() or body.ward.strip()
    _ward = body.ward.strip() or (ward_master.ward_for_colony(_colony) or "")
    doc = {
        "id": str(uuid.uuid4()), "serial_number": None, "property_id": pid,
        "owner_name": body.owner_name.strip(), "mobile": body.mobile.strip(),
        "alternate_mobile": body.alternate_mobile.strip(), "address": body.address.strip(),
        "colony": _colony, "ward": _ward,
        "latitude": body.latitude, "longitude": body.longitude, "category": body.category,
        "amount": "0", "status": "Pending", "source": "field",
        "assigned_employee_id": user["id"], "assigned_employee_name": user.get("name"),
        "created_by": user["id"], "created_at": ts,
    }
    await db.properties.insert_one(doc)
    await clear_map_cache()
    await audit(user, "FIELD_PROPERTY_CREATE", "property", doc["id"], meta={"property_id": pid})
    doc.pop("_id", None)
    return {"id": doc["id"], "property_id": pid, "property": doc}



@phed_router.post("/surveys/{survey_id}/approve")
async def approve_survey(survey_id: str, user: dict = Depends(get_current_user)):
    require_officer(user)
    db = get_db()
    survey = await db.phed_surveys.find_one({"id": survey_id}, {"_id": 0})
    if not survey:
        raise HTTPException(404, "Survey not found")
    if survey["status"] not in ("Submitted", "Requires Review", "Document Pending"):
        raise HTTPException(409, f"Survey is {survey['status']}; only Submitted / Requires Review / Document Pending surveys can be approved")
    await _approve_survey_doc(db, survey, user)
    await clear_map_cache()
    return {"message": "Survey approved", "status": "Approved"}


async def _approve_survey_doc(db, survey: dict, user: dict):
    """Core approve logic (shared by single + bulk approve)."""
    survey_id = survey["id"]
    ts = now_iso()
    await db.phed_surveys.update_one({"id": survey_id}, {"$set": {"status": "Approved", "approved_at": ts, "approved_by": user["id"], "approved_by_name": user.get("name"), "updated_at": ts}})
    phed_status = "No PHED Connection" if survey.get("survey_type") == "NO_CONNECTION" else "Approved"
    await db.properties.update_one({"id": survey["property_record_id"]}, {"$set": {"phed_survey_status": phed_status, "phed_survey_state": "Approved", "phed_outcome": _survey_outcome(survey)}})
    await audit(user, "SURVEY_APPROVE", "survey", survey_id, before={"status": survey["status"]}, after={"status": "Approved"})
    await record_audit(user, "PHED_SURVEY_APPROVE", "phed_survey", survey_id,
                       {"property_record_id": survey.get("property_record_id"), "before": {"status": survey["status"]}, "after": {"status": "Approved"}})


class BulkApproveBody(BaseModel):
    ids: List[str] = []
    all_pending: bool = False


@phed_router.post("/surveys/bulk-approve")
async def bulk_approve_surveys(body: BulkApproveBody, user: dict = Depends(get_current_user)):
    """Approve many yellow surveys at once (Approval Queue → select & approve all)."""
    require_officer(user)
    db = get_db()
    q: Dict[str, Any] = {"status": {"$in": ["Submitted", "Requires Review", "Document Pending"]}}
    if not body.all_pending:
        if not body.ids:
            raise HTTPException(400, "No surveys selected")
        q["id"] = {"$in": body.ids}
    surveys = await db.phed_surveys.find(q, {"_id": 0}).to_list(None)
    for s in surveys:
        await _approve_survey_doc(db, s, user)
    await clear_map_cache()
    return {"message": f"{len(surveys)} survey(s) approved", "approved": len(surveys)}


# ---------------- office documents → auto-approve from Excel ----------------
OFFICE_HEADERS = {"property_id": ["property id", "pid", "property id pid", "mc property id", "property no"],
                  "consumer_id": ["consumer id", "consumer no", "phed consumer id", "consumer number"],
                  "owner_name": ["owner name", "name", "consumer name"], "mobile": ["mobile", "phone", "phone no", "mobile no"],
                  "remarks": ["remarks", "remark", "note"]}


@phed_router.get("/office-import/sample")
async def office_import_sample(user: dict = Depends(get_current_user)):
    require_admin(user)
    wb = Workbook(); ws = wb.active; ws.title = "Office Documents"
    ws.append(["Property ID (PID)", "Consumer ID", "Owner Name", "Mobile", "Remarks"])
    ws.append(["THS-0001", "", "Ram Kumar", "9876543210", "Documents received at office"])
    ws.append(["", "4326479", "Suman", "7206864297", "Bill copy submitted"])
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": 'attachment; filename="phed_office_documents_sample.xlsx"'})


@phed_router.post("/office-import")
async def office_import(file: UploadFile = File(...), user: dict = Depends(get_current_user)):
    """Admin uploads an Excel of properties/consumers whose documents were received at the office.
    Each matched property gets an Approved (locked) survey with source 'office'."""
    require_admin(user)
    if not (file.filename or "").lower().endswith((".xlsx", ".xlsm")):
        raise HTTPException(400, "Only .xlsx files are supported")
    content = await file.read()
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise HTTPException(400, "Empty sheet")
    headers = [norm_header(h) for h in rows[0]]
    col = {}
    for field, aliases in OFFICE_HEADERS.items():
        for i, h in enumerate(headers):
            if h in aliases:
                col[field] = i; break
    if "property_id" not in col and "consumer_id" not in col:
        raise HTTPException(400, f"Need a 'Property ID (PID)' or 'Consumer ID' column. Found: {', '.join(str(h) for h in rows[0])}")
    db = get_db()
    ts = now_iso()
    res = dict(total_rows=0, approved=0, already_approved=0, not_found=0, invalid=0)
    unmatched = []
    get = lambda r, f: cell_str(r[col[f]]) if f in col and col[f] < len(r) else ""  # noqa: E731
    batch_id = str(uuid.uuid4())
    for idx, r in enumerate(rows[1:], start=2):
        if not r or all(v in (None, "") for v in r):
            continue
        res["total_rows"] += 1
        pid, cid = get(r, "property_id"), get(r, "consumer_id")
        if not pid and not cid:
            res["invalid"] += 1; unmatched.append({"row": idx, "reason": "No Property ID or Consumer ID"}); continue
        prop = None
        consumer = None
        if pid:
            prop = await db.properties.find_one({"property_id": {"$regex": f"^{re.escape(pid)}$", "$options": "i"}}, {"_id": 0})
        if cid:
            consumer = await db.phed_consumers.find_one({"consumer_id_norm": norm_id(cid)}, {"_id": 0})
            if not prop and consumer and consumer.get("linked_property_id"):
                prop = await db.properties.find_one({"id": consumer["linked_property_id"]}, {"_id": 0})
        if not prop and consumer:
            # Auto-link: only Consumer ID given (no PID) and consumer not yet linked to a property.
            # 1) match by consumer phone == property mobile, 2) by owner/consumer name + colony.
            ph = norm_id(consumer.get("phone") or "")
            if ph and len(ph) >= 7 and len(set(ph)) > 1:
                prop = await db.properties.find_one({"$or": [{"mobile": consumer.get("phone")}, {"alternate_mobile": consumer.get("phone")}]}, {"_id": 0})
            if not prop and consumer.get("consumer_name"):
                nm = {"$regex": f"^{re.escape(consumer['consumer_name'].strip())}$", "$options": "i"}
                pq2: Dict[str, Any] = {"owner_name": nm}
                if consumer.get("colony_name") or consumer.get("locality"):
                    pq2["colony"] = {"$regex": f"^{re.escape((consumer.get('colony_name') or consumer.get('locality')).strip())}$", "$options": "i"}
                prop = await db.properties.find_one(pq2, {"_id": 0})
            if not prop:
                # 3) still nothing — create a property record from the consumer so it appears & can be approved.
                new_pid = f"PHED-{norm_id(cid).upper()}" if cid else f"PHED-{str(uuid.uuid4())[:8].upper()}"
                prop = {"id": str(uuid.uuid4()), "property_id": new_pid, "serial_number": None,
                        "owner_name": consumer.get("consumer_name") or get(r, "owner_name") or "", "mobile": consumer.get("phone") or "",
                        "address": consumer.get("address") or "", "colony": consumer.get("colony_name") or consumer.get("locality") or "",
                        "ward": consumer.get("ward_number") or "", "latitude": None, "longitude": None, "category": consumer.get("category") or "Domestic",
                        "status": "Pending", "source": "office", "created_by": user["id"], "created_at": ts}
                await db.properties.insert_one(dict(prop))
                res["created_properties"] = res.get("created_properties", 0) + 1
            # link the consumer to the resolved property
            await db.phed_consumers.update_one({"id": consumer["id"]}, {"$set": {"linked_property_id": prop["id"], "linked_property_number": prop.get("property_id"), "updated_at": ts}})
            res["auto_linked"] = res.get("auto_linked", 0) + 1
        if not prop:
            res["not_found"] += 1
            unmatched.append({"row": idx, "property_id": pid, "consumer_id": cid, "reason": "Consumer not found" if cid and not consumer else "Property not found"})
            continue
        existing = await db.phed_surveys.find_one({"property_record_id": prop["id"], "status": {"$ne": "Rejected"}}, {"_id": 0})
        if existing and existing["status"] == "Approved":
            res["already_approved"] += 1; continue
        water = {"has_connection": True, "consumer_ref": consumer["id"] if consumer else None, "consumer_id": cid or (consumer or {}).get("consumer_id", ""),
                 "consumer_name": get(r, "owner_name") or (consumer or {}).get("consumer_name") or prop.get("owner_name") or "",
                 "phone": get(r, "mobile") or (consumer or {}).get("phone") or prop.get("mobile") or "", "mobile": get(r, "mobile") or prop.get("mobile") or "",
                 "office_documents": True, "document_pending": False, "missing_documents": []}
        upd = {"status": "Approved", "source": "office", "office_batch_id": batch_id, "survey_type": "WATER_CONNECTION", "water": water,
               "remarks": get(r, "remarks") or None, "consumer_refs": [consumer["id"]] if consumer else [], "consumer_ref": consumer["id"] if consumer else None,
               "submitted_at": ts, "approved_at": ts, "approved_by": user["id"], "approved_by_name": user.get("name"), "updated_at": ts}
        if existing:
            await db.phed_surveys.update_one({"id": existing["id"]}, {"$set": upd})
            sid = existing["id"]
        else:
            sid = str(uuid.uuid4())
            ref = await _next_reference(db, "OF")
            await db.phed_surveys.insert_one({"id": sid, "property_record_id": prop["id"], "property_id": prop["property_id"], "surveyor_id": user["id"], "surveyor_name": user.get("name"),
                                              "ward_number": prop.get("ward"), "colony_name": prop.get("colony"), "attachments": [], "started_at": ts, "created_at": ts,
                                              "latitude": prop.get("latitude"), "longitude": prop.get("longitude"), "reference_number": ref, "reference_code": "OF", **upd})
        await db.properties.update_one({"id": prop["id"]}, {"$set": {"phed_survey_status": "Approved", "phed_survey_type": "WATER_CONNECTION", "phed_surveyed_at": ts,
                                                                     "phed_survey_state": "Approved", "phed_outcome": "HAS_CONNECTION", "phed_survey_source": "office"}})
        if consumer:
            await db.phed_consumers.update_one({"id": consumer["id"]}, {"$set": {"status": "Surveyed", "linked_property_id": prop["id"], "linked_property_number": prop["property_id"], "updated_at": ts}})
        res["approved"] += 1
    await clear_map_cache()
    await audit(user, "OFFICE_IMPORT", "import", batch_id, meta={**res, "filename": file.filename})
    await db.phed_office_imports.insert_one({"id": batch_id, "filename": file.filename, "uploaded_by": user["id"], "uploaded_by_name": user.get("name"),
                                             "result": res, "unmatched": unmatched[:500], "created_at": ts})
    return {"batch_id": batch_id, **res, "unmatched": unmatched[:200]}


@phed_router.get("/office-imports")
async def list_office_imports(user: dict = Depends(get_current_user)):
    require_admin(user)
    return await get_db().phed_office_imports.find({}, {"_id": 0, "unmatched": 0}).sort("created_at", -1).limit(50).to_list(None)


@phed_router.get("/surveys/mine")
async def my_surveys(status: Optional[str] = None, page: int = 1, limit: int = 30, user: dict = Depends(get_current_user)):
    q: Dict[str, Any] = {"surveyor_id": user["id"]}
    if status:
        q["status"] = status
    total = await get_db().phed_surveys.count_documents(q)
    items = await get_db().phed_surveys.find(q, {"_id": 0}).sort("updated_at", -1).skip((page - 1) * limit).limit(limit).to_list(None)
    return {"surveys": [survey_public(s) for s in items], "total": total}


@phed_router.get("/surveys")
async def list_surveys(status: Optional[str] = None, surveyor_id: Optional[str] = None, ward_id: Optional[str] = None, colony: Optional[str] = None,
                       search: Optional[str] = None, queue: Optional[str] = None, page: int = 1, limit: int = 30, user: dict = Depends(get_current_user)):
    require_officer(user)
    q: Dict[str, Any] = {}
    if queue == "pending":
        q["status"] = {"$in": ["Submitted", "Requires Review", "Document Pending"]}
    elif status:
        q["status"] = status
    if surveyor_id: q["surveyor_id"] = surveyor_id
    if ward_id: q["ward_id"] = ward_id
    if colony: q["colony_name"] = colony
    if search and search.strip():
        rx = {"$regex": re.escape(search.strip()), "$options": "i"}
        q["$or"] = [{"reference_number": rx}, {"property_id": rx}, {"water.consumer_id": rx},
                    {"water.consumer_name": rx}, {"water.new_owner_name": rx}, {"water.connection_numbers": rx}]
    total = await get_db().phed_surveys.count_documents(q)
    items = await get_db().phed_surveys.find(q, {"_id": 0}).sort("updated_at", -1).skip((page - 1) * limit).limit(limit).to_list(None)
    return {"surveys": [survey_public(s) for s in items], "total": total, "pages": (total + limit - 1) // limit}


@phed_router.get("/map-summary")
async def map_summary(colony: Optional[str] = None, user: dict = Depends(get_current_user)):
    """Live red/yellow/green survey counts for the whole town (or one colony) —
    shown on the map legend so admins/surveyors see progress without choosing an area."""
    db = get_db()
    base: Dict[str, Any] = {}
    if not is_officer(user):
        base["$or"] = [{"assigned_employee_id": user["id"]}, {"assigned_employee_ids": user["id"]}]
    if colony and colony.strip():
        base["colony"] = {"$regex": f"^{re.escape(colony.strip())}$", "$options": "i"}
    green_cond = [{"phed_survey_state": "Approved"}, {"phed_survey_status": {"$in": ["Approved", "No PHED Connection"]}}]
    yellow_states = ["Submitted", "Requires Review", "Document Pending", "Draft"]
    total = await db.properties.count_documents(base)
    green = await db.properties.count_documents({**base, "$and": [{"$or": green_cond}]})
    yellow = await db.properties.count_documents({**base, "phed_survey_state": {"$in": yellow_states}})
    red = max(total - green - yellow, 0)
    return {"total": total, "green": green, "yellow": yellow, "red": red}


@phed_router.get("/location-pending")
async def location_pending(user: dict = Depends(get_current_user)):
    """Office-import created properties that still have no GPS point — admin sets location on the map later."""
    require_officer(user)
    db = get_db()
    q = {"source": "office", "$or": [{"latitude": None}, {"latitude": {"$exists": False}}, {"longitude": None}, {"longitude": {"$exists": False}}]}
    rows = await db.properties.find(q, {"_id": 0, "id": 1, "property_id": 1, "owner_name": 1, "mobile": 1, "address": 1, "colony": 1, "ward": 1, "phed_survey_status": 1, "created_at": 1}).sort("created_at", -1).limit(1000).to_list(1000)
    return {"properties": rows, "total": len(rows)}


class SetLocationBody(BaseModel):
    latitude: float
    longitude: float


@phed_router.post("/properties/{property_id}/location")
async def set_property_location(property_id: str, body: SetLocationBody, user: dict = Depends(get_current_user)):
    """Set/adjust a property's GPS point (used by the Location Pending list)."""
    require_officer(user)
    db = get_db()
    prop = await db.properties.find_one({"id": property_id}, {"_id": 0, "id": 1})
    if not prop:
        raise HTTPException(404, "Property not found")
    await db.properties.update_one({"id": property_id}, {"$set": {"latitude": body.latitude, "longitude": body.longitude, "location_set_by": user["id"], "location_set_at": now_iso()}})
    await clear_map_cache()
    await audit(user, "PROPERTY_LOCATION_SET", "property", property_id, after={"latitude": body.latitude, "longitude": body.longitude})
    return {"message": "Location saved", "latitude": body.latitude, "longitude": body.longitude}


class BulkLocateBody(BaseModel):
    ids: List[str] = []
    all_pending: bool = False


@phed_router.post("/location-pending/bulk-locate")
async def bulk_locate(body: BulkLocateBody, user: dict = Depends(get_current_user)):
    """Auto-place office properties at their colony's centre (mean GPS of already-located
    properties in the same colony). Skips a property if its colony has no reference point yet."""
    require_officer(user)
    db = get_db()
    q: Dict[str, Any] = {"source": "office", "$or": [{"latitude": None}, {"latitude": {"$exists": False}}, {"longitude": None}, {"longitude": {"$exists": False}}]}
    if not body.all_pending:
        if not body.ids:
            raise HTTPException(400, "No properties selected")
        q["id"] = {"$in": body.ids}
    targets = await db.properties.find(q, {"_id": 0, "id": 1, "colony": 1, "ward": 1}).to_list(None)
    ts = now_iso()
    centre_cache: Dict[str, Any] = {}

    async def colony_centre(colony: str):
        key = (colony or "").strip().lower()
        if key in centre_cache:
            return centre_cache[key]
        rows = await db.properties.aggregate([
            {"$match": {"colony": {"$regex": f"^{re.escape((colony or '').strip())}$", "$options": "i"},
                        "latitude": {"$ne": None, "$exists": True}, "longitude": {"$ne": None, "$exists": True}}},
            {"$group": {"_id": None, "lat": {"$avg": "$latitude"}, "lng": {"$avg": "$longitude"}, "n": {"$sum": 1}}},
        ]).to_list(1)
        centre = (round(rows[0]["lat"], 6), round(rows[0]["lng"], 6)) if rows and rows[0].get("n") else None
        centre_cache[key] = centre
        return centre

    located, skipped = 0, []
    import random
    for t in targets:
        centre = await colony_centre(t.get("colony"))
        if not centre:
            skipped.append({"id": t["id"], "colony": t.get("colony"), "reason": "colony me koi located property nahi"})
            continue
        # tiny jitter so multiple auto-placed pins don't stack exactly on top of each other
        lat = round(centre[0] + random.uniform(-0.00025, 0.00025), 6)
        lng = round(centre[1] + random.uniform(-0.00025, 0.00025), 6)
        await db.properties.update_one({"id": t["id"]}, {"$set": {"latitude": lat, "longitude": lng, "location_set_by": user["id"], "location_set_at": ts, "location_auto": True}})
        located += 1
    if located:
        await clear_map_cache()
    return {"message": f"{located} property colony-centre par place ho gayi", "located": located, "skipped": skipped}


@phed_router.post("/surveys/{survey_id}/reject")
async def reject_survey(survey_id: str, reason: str = Form(...), user: dict = Depends(get_current_user)):
    require_admin(user)
    db = get_db()
    survey = await db.phed_surveys.find_one({"id": survey_id}, {"_id": 0})
    if not survey:
        raise HTTPException(404, "Survey not found")
    ts = now_iso()
    await db.phed_surveys.update_one({"id": survey_id}, {"$set": {"status": "Rejected", "rejection_reason": reason, "updated_at": ts}})
    await db.properties.update_one({"id": survey["property_record_id"]}, {"$set": {"phed_survey_status": "Rejected", "phed_survey_state": "Rejected", "phed_outcome": None}})
    await clear_map_cache()
    await audit(user, "SURVEY_REJECT", "survey", survey_id, before={"status": survey["status"]}, after={"status": "Rejected", "reason": reason})
    await record_audit(user, "PHED_SURVEY_REJECT", "phed_survey", survey_id,
                       {"property_record_id": survey.get("property_record_id"), "before": {"status": survey["status"]},
                        "after": {"status": "Rejected"}, "reason": reason})
    return {"message": "Survey rejected; property reopened for re-survey"}


@phed_router.post("/surveys/{survey_id}/reopen")
async def reopen_survey(survey_id: str, reason: str = Form(""), mode: str = Form("correction"), user: dict = Depends(get_current_user)):
    """Admin/officer sends a survey back to the surveyor as Document Pending, keeping all
    entered data and uploaded documents. mode='correction' (return for correction) or
    mode='pending' (mark pending). The surveyor can reopen and complete it."""
    require_officer(user)
    db = get_db()
    survey = await db.phed_surveys.find_one({"id": survey_id}, {"_id": 0})
    if not survey:
        raise HTTPException(404, "Survey not found")
    if survey["status"] in ("Draft", "Rejected"):
        raise HTTPException(409, f"Survey is {survey['status']}; nothing to return")
    ts = now_iso()
    note = (reason or "").strip() or None
    action = "Marked pending" if mode == "pending" else "Returned for correction"
    water = dict(survey.get("water") or {})
    water["document_pending"] = True
    if note:
        water["return_reason"] = note
    upd = {"status": "Document Pending", "return_reason": note, "return_mode": mode,
           "returned_by": user["id"], "returned_by_name": user.get("name"), "returned_at": ts,
           "water": water, "approved_at": None, "approved_by": None, "approved_by_name": None, "updated_at": ts}
    await db.phed_surveys.update_one({"id": survey_id}, {"$set": upd})
    await db.properties.update_one({"id": survey["property_record_id"]}, {"$set": {"phed_survey_status": "Document Pending", "phed_survey_state": "Document Pending"}})
    await clear_map_cache()
    await audit(user, "SURVEY_RETURN", "survey", survey_id, before={"status": survey["status"]},
                after={"status": "Document Pending", "mode": mode, "reason": note})
    await record_audit(user, "PHED_SURVEY_RETURN", "phed_survey", survey_id,
                       {"property_record_id": survey.get("property_record_id"), "before": {"status": survey["status"]},
                        "after": {"status": "Document Pending"}, "mode": mode, "reason": note})
    return {"message": f"{action}; surveyor can complete it", "status": "Document Pending"}


# ---------------- dashboard / audit / export ----------------
def _survey_filter(ward_id, colony, surveyor_id, status, date_from, date_to):
    q: Dict[str, Any] = {}
    if ward_id: q["ward_id"] = ward_id
    if colony: q["colony_name"] = colony
    if surveyor_id: q["surveyor_id"] = surveyor_id
    if status: q["status"] = status
    if date_from or date_to:
        rng = {}
        if date_from: rng["$gte"] = date_from
        if date_to: rng["$lte"] = date_to + "T23:59:59"
        q["updated_at"] = rng
    return q


@phed_router.get("/my-progress")
async def my_phed_progress(user: dict = Depends(get_current_user)):
    db = get_db()
    uid = user["id"]
    assigned_q = {"$or": [{"assigned_employee_id": uid}, {"assigned_employee_ids": uid}]}
    total_props = await db.properties.count_documents(assigned_q)
    surveyed_prop_ids = set()
    completed = in_progress = no_conn = new_unlisted = 0
    total_submitted = already_connection = sewer_connection = new_connection = ownership_change = death_transfer = 0
    async for s in db.phed_surveys.find(
        {"surveyor_id": uid},
        {"_id": 0, "property_record_id": 1, "status": 1, "survey_type": 1, "water": 1},
    ):
        surveyed_prop_ids.add(s.get("property_record_id"))
        st = s.get("status")
        if st in OPEN_STATUSES or st in ("Completed", "No Connection"):
            completed += 1
            if s.get("survey_type") == "NO_CONNECTION" or st == "No Connection":
                no_conn += 1
            if s.get("survey_type") == "NEW_UNLISTED":
                new_unlisted += 1
        elif st in ("In Progress", "Draft"):
            in_progress += 1
        # Outcome breakdown over everything the surveyor has actually submitted (not drafts)
        if st and st != "Draft":
            total_submitted += 1
            w = s.get("water") or {}
            if w.get("has_connection"):
                already_connection += 1
            if w.get("has_sewer") or (w.get("sewer_connection_numbers") or []):
                sewer_connection += 1
            if w.get("new_connection") or s.get("survey_type") == "NO_CONNECTION":
                new_connection += 1
            oc = w.get("owner_change")
            if oc == "OWNERSHIP_CHANGE":
                ownership_change += 1
            elif oc == "DEATH_TRANSFER":
                death_transfer += 1
    done_props = len([p for p in surveyed_prop_ids if p])
    pending = max(total_props - done_props, 0)
    field_properties = await db.properties.count_documents({"source": "field", "created_by": uid})
    return {
        "assigned_area": user.get("assigned_area"),
        "total_properties": total_props,
        "phed_pending": pending,
        "phed_in_progress": in_progress,
        "phed_completed": completed,
        "no_connection": no_conn,
        "new_unlisted": new_unlisted,
        "total_submitted": total_submitted,
        "already_connection": already_connection,
        "sewer_connection": sewer_connection,
        "new_connection": new_connection,
        "ownership_change": ownership_change,
        "death_transfer": death_transfer,
        "field_properties": field_properties,
    }



@phed_router.get("/dashboard")
async def phed_dashboard(ward_id: Optional[str] = None, colony: Optional[str] = None, surveyor_id: Optional[str] = None, status: Optional[str] = None,
                         service: Optional[str] = None, category: Optional[str] = None, source: Optional[str] = None, linked: Optional[str] = None,
                         date_from: Optional[str] = None, date_to: Optional[str] = None,
                         user: dict = Depends(get_current_user)):
    require_officer(user)
    db = get_db()
    cq: Dict[str, Any] = {"is_active": True}
    if ward_id: cq["ward_id"] = ward_id
    if colony: cq["colony_name"] = colony
    if category: cq["category"] = category
    if source in ("imported", "survey"): cq["source"] = source
    if linked == "yes": cq["linked_property_id"] = {"$ne": None}
    if linked == "no": cq["linked_property_id"] = None
    consumer_refs = None
    if ward_id or colony or category or source or linked:
        consumer_refs = await db.phed_consumers.distinct("id", cq)
    conq: Dict[str, Any] = {"is_active": True}
    if consumer_refs is not None: conq["consumer_ref"] = {"$in": consumer_refs}
    if service in SERVICES: conq["service"] = service
    total_consumers = await db.phed_consumers.count_documents(cq)
    pipeline = [{"$match": conq}, {"$group": {"_id": "$service", "n": {"$sum": 1}, "consumers": {"$addToSet": "$consumer_ref"}}}]
    by_service = {r["_id"]: r for r in await db.phed_connections.aggregate(pipeline).to_list(None)}
    water_set = set(by_service.get("Water", {}).get("consumers", []))
    sewer_set = set(by_service.get("Sewer", {}).get("consumers", []))
    sq = _survey_filter(ward_id, colony, surveyor_id, status, date_from, date_to)
    if service in SERVICES:
        sq["consumer_refs"] = {"$in": list(water_set if service == "Water" else sewer_set)}
    if category:
        sq["consumer_refs"] = {"$in": consumer_refs}
    if source or linked:
        sq["consumer_refs"] = {"$in": consumer_refs}
    s_pipeline = [{"$match": sq}, {"$group": {"_id": "$status", "n": {"$sum": 1}}}]
    by_status = {r["_id"]: r["n"] for r in await db.phed_surveys.aggregate(s_pipeline).to_list(None)}
    # Property universe for the selected ward/colony (target = every existing MC property)
    pq: Dict[str, Any] = {}
    if colony:
        pq["colony"] = colony
    elif ward_id:
        w = await db.phed_wards.find_one({"id": ward_id}, {"_id": 0, "colonies": 1})
        pq["colony"] = {"$in": (w or {}).get("colonies", [])}
    total_properties = await db.properties.count_documents(pq)
    props_surveyed = await db.properties.count_documents({**pq, "phed_survey_status": {"$in": ["Submitted", "Requires Review", "Approved", "No PHED Connection"]}})
    props_in_progress = await db.properties.count_documents({**pq, "phed_survey_status": "Draft"})
    props_no_conn = await db.properties.count_documents({**pq, "phed_survey_status": "No PHED Connection"})
    props_linked = len(await db.phed_consumers.distinct("linked_property_id", {"is_active": True, "linked_property_id": {"$ne": None}}))
    multi_conn = len([r for r in await db.phed_connections.aggregate([{"$match": {"is_active": True, "linked_property_id": {"$ne": None}}},
                                                                     {"$group": {"_id": "$linked_property_id", "n": {"$sum": 1}}}, {"$match": {"n": {"$gt": 1}}}]).to_list(None)])
    colonies_all = [c for c in await db.properties.distinct("colony") if c]
    if ward_id:
        w = await db.phed_wards.find_one({"id": ward_id}, {"_id": 0, "colonies": 1})
        colonies_all = w.get("colonies", []) if w else []
    if colony:
        colonies_all = [colony]
    completed_colonies = 0
    for cname in colonies_all:
        total_p = await db.properties.count_documents({"colony": cname})
        if total_p and total_p == await db.properties.count_documents({"colony": cname, "phed_survey_status": {"$in": ["Approved", "No PHED Connection"]}}):
            completed_colonies += 1
    completed_wards = 0
    for w in await db.phed_wards.find({} if not ward_id else {"id": ward_id}, {"_id": 0, "colonies": 1}).to_list(None):
        wq = {"colony": {"$in": w.get("colonies", [])}}
        total_p = await db.properties.count_documents(wq)
        if total_p and total_p == await db.properties.count_documents({**wq, "phed_survey_status": {"$in": ["Approved", "No PHED Connection"]}}):
            completed_wards += 1
    active_surveyors = len(await db.phed_surveys.distinct("surveyor_id", {**sq, "status": {"$in": ["Draft", "Submitted", "Requires Review", "Approved"]}}))
    linked_c = await db.phed_consumers.count_documents({**cq, "linked_property_id": {"$ne": None}})
    unlisted = await db.phed_consumers.count_documents({**cq, "source": "survey"})
    unlisted_conns = await db.phed_connections.count_documents({"is_active": True, "source": "survey"})
    type_pipeline = [{"$match": {**sq, "status": {"$in": OPEN_STATUSES}}}, {"$group": {"_id": "$survey_type", "n": {"$sum": 1}}}]
    by_type = {r["_id"]: r["n"] for r in await db.phed_surveys.aggregate(type_pipeline).to_list(None)}
    # Survey OUTCOME breakdown (locked / denied / new / has-connection) — PHED, not house-tax
    outcome_rows = await db.phed_surveys.aggregate([
        {"$match": {**sq, "status": {"$in": ["Submitted", "Requires Review", "Document Pending", "Approved"]}}},
        {"$project": {"oc": {"$switch": {"branches": [
            {"case": {"$eq": ["$water.property_locked", True]}, "then": "LOCKED"},
            {"case": {"$eq": ["$water.owner_denied", True]}, "then": "DENIED"},
            {"case": {"$eq": ["$survey_type", "NO_CONNECTION"]}, "then": "NEW"},
            {"case": {"$eq": ["$water.new_connection", True]}, "then": "NEW"},
            {"case": {"$eq": ["$water.has_connection", True]}, "then": "HAS_CONNECTION"},
        ], "default": "OTHER"}}}},
        {"$group": {"_id": "$oc", "n": {"$sum": 1}}},
    ]).to_list(None)
    by_outcome = {r["_id"]: r["n"] for r in outcome_rows}
    # Surveyor-wise leaderboard (submitted + approved by each surveyor)
    surveyor_rows = await db.phed_surveys.aggregate([
        {"$match": {**sq, "status": {"$in": ["Submitted", "Requires Review", "Document Pending", "Approved"]}}},
        {"$group": {"_id": {"id": "$surveyor_id", "name": "$surveyor_name"},
                    "total": {"$sum": 1},
                    "approved": {"$sum": {"$cond": [{"$eq": ["$status", "Approved"]}, 1, 0]}}}},
        {"$sort": {"total": -1}}, {"$limit": 50},
    ]).to_list(None)
    # Today's submissions per surveyor (for daily target vs done)
    today_iso = now_iso()[:10]
    today_rows = await db.phed_surveys.aggregate([
        {"$match": {**sq, "submitted_at": {"$gte": today_iso}}},
        {"$group": {"_id": "$surveyor_id", "n": {"$sum": 1}}},
    ]).to_list(None)
    today_map = {r["_id"]: r["n"] for r in today_rows}
    by_surveyor = [{"id": r["_id"].get("id"), "name": r["_id"].get("name") or "—",
                    "total": r["total"], "approved": r["approved"], "pending": r["total"] - r["approved"],
                    "today": today_map.get(r["_id"].get("id"), 0), "target": DAILY_TARGET_DEFAULT}
                   for r in surveyor_rows if r["_id"].get("id")]
    return {
        "total_properties": total_properties, "target_properties": total_properties,
        "properties_pending": max(total_properties - props_surveyed - props_in_progress, 0), "properties_in_progress": props_in_progress,
        "properties_completed": props_surveyed, "properties_linked": props_linked, "properties_not_linked": max(total_properties - props_linked, 0),
        "properties_no_connection": props_no_conn, "properties_multi_connection": multi_conn, "new_unlisted_connections": unlisted_conns,
        "total_consumers": total_consumers, "linked_consumers": linked_c, "survey_created_consumers": unlisted,
        "total_connections": sum(r["n"] for r in by_service.values()), "water_connections": by_service.get("Water", {}).get("n", 0),
        "sewer_connections": by_service.get("Sewer", {}).get("n", 0), "both_water_sewer": len(water_set & sewer_set),
        "surveys_total": sum(by_status.values()), "pending_surveys": by_status.get("Draft", 0) + by_status.get("Not Started", 0),
        "submitted_surveys": by_status.get("Submitted", 0), "review_surveys": by_status.get("Requires Review", 0),
        "approved_surveys": by_status.get("Approved", 0), "rejected_surveys": by_status.get("Rejected", 0),
        "total_colonies": len(colonies_all), "completed_colonies": completed_colonies, "completed_wards": completed_wards, "active_surveyors": active_surveyors,
        "by_survey_type": {"existing_linked": by_type.get("EXISTING_LINKED", 0), "water_connection": by_type.get("WATER_CONNECTION", 0), "new_unlisted": by_type.get("NEW_UNLISTED", 0), "no_connection": by_type.get("NO_CONNECTION", 0)},
        "by_outcome": {"new_connection": by_outcome.get("NEW", 0), "property_locked": by_outcome.get("LOCKED", 0),
                       "owner_denied": by_outcome.get("DENIED", 0), "has_connection": by_outcome.get("HAS_CONNECTION", 0)},
        "by_surveyor": by_surveyor,
        "gps_flagged": await db.phed_surveys.count_documents({**sq, "gps_flagged": True}),
    }


@phed_router.get("/filters")
async def filter_options(user: dict = Depends(get_current_user)):
    require_officer(user)
    db = get_db()
    wards = await db.phed_wards.find({}, {"_id": 0, "id": 1, "ward_number": 1, "name": 1}).sort("ward_number", 1).to_list(None)
    colonies = sorted(c for c in await db.properties.distinct("colony") if c)
    surveyor_ids = await db.phed_surveys.distinct("surveyor_id")
    surveyors = await db.phed_surveys.aggregate([{"$group": {"_id": "$surveyor_id", "name": {"$first": "$surveyor_name"}}}]).to_list(None)
    return {"wards": wards, "colonies": colonies, "surveyors": [{"id": s["_id"], "name": s["name"]} for s in surveyors if s["_id"] in surveyor_ids],
            "services": SERVICES, "categories": CATEGORIES, "statuses": ["Draft", "Submitted", "Requires Review", "Document Pending", "Approved", "Rejected"],
            "sources": [["imported", "Existing (imported)"], ["survey", "New / unlisted"]], "linked": [["yes", "Linked to property"], ["no", "Not linked"]]}


@phed_router.get("/audit")
async def audit_list(entity_id: Optional[str] = None, action: Optional[str] = None, page: int = 1, limit: int = 50, user: dict = Depends(get_current_user)):
    require_officer(user)
    q: Dict[str, Any] = {}
    if entity_id: q["entity_id"] = entity_id
    if action: q["action"] = action
    total = await get_db().phed_audit_logs.count_documents(q)
    items = await get_db().phed_audit_logs.find(q, {"_id": 0}).sort("timestamp", -1).skip((page - 1) * limit).limit(limit).to_list(None)
    return {"logs": items, "total": total}


@phed_router.get("/export")
async def export_phed(ward_id: Optional[str] = None, colony: Optional[str] = None, surveyor_id: Optional[str] = None, status: Optional[str] = None,
                      service: Optional[str] = None, category: Optional[str] = None, source: Optional[str] = None, date_from: Optional[str] = None, date_to: Optional[str] = None,
                      user: dict = Depends(get_current_user)):
    require_officer(user)
    if user["role"] not in ("ADMIN", "MC_OFFICER") and "export" not in (user.get("permissions") or []):
        raise HTTPException(403, "Export permission required")
    db = get_db()
    cq: Dict[str, Any] = {"is_active": True}
    if ward_id: cq["ward_id"] = ward_id
    if colony: cq["colony_name"] = colony
    if category: cq["category"] = category
    if source: cq["source"] = source
    if service in SERVICES:
        cq["id"] = {"$in": await db.phed_connections.distinct("consumer_ref", {"service": service, "is_active": True})}
    consumers = await db.phed_consumers.find(cq, {"_id": 0, "original_values": 0}).sort([("ward_number", 1), ("consumer_id", 1)]).to_list(None)
    await attach_connections(consumers)
    prop_ids = list({c["linked_property_id"] for c in consumers if c.get("linked_property_id")})
    props = {p["id"]: p for p in await db.properties.find({"id": {"$in": prop_ids}}, {"_id": 0, "id": 1, "property_id": 1, "owner_name": 1, "address": 1, "colony": 1, "ward": 1, "latitude": 1, "longitude": 1}).to_list(None)}
    sq = _survey_filter(None, None, surveyor_id, status, date_from, date_to)
    sq["property_record_id"] = {"$in": prop_ids}
    surveys = {s["property_record_id"]: s for s in await db.phed_surveys.find(sq, {"_id": 0}).sort("updated_at", 1).to_list(None)}
    if surveyor_id or status or date_from or date_to:
        consumers = [c for c in consumers if c.get("linked_property_id") in surveys]
    wards = {w["id"]: w for w in await db.phed_wards.find({}, {"_id": 0}).to_list(None)}
    town = await db.properties.find_one({}, {"_id": 0, "town": 1}) or {}

    wb = Workbook()
    ws = wb.active
    ws.title = "PHED Survey"
    headers = ["Town", "Ward", "Colony", "Property ID", "Property Owner", "Property Address", "Property Latitude", "Property Longitude",
               "PHED Consumer ID", "Consumer Name", "Father/Husband Name", "PPP ID", "Consumer Address", "Locality", "Phone Number",
               "Water Connection Numbers", "Sewer Connection Numbers", "Water Connections", "Sewer Connections", "Total Connections", "Connection Category",
               "Existing/New/No Connection", "Property Linking Status", "Name Mismatch", "Survey Status", "Surveyor",
               "Survey Latitude", "Survey Longitude", "GPS Accuracy (m)", "Survey Date", "Document Availability", "Remarks", "Import Source"]
    ws.append(headers)
    from server import get_current_town_code
    town_code = get_current_town_code()
    type_label = {"EXISTING_LINKED": "Existing", "WATER_CONNECTION": "Water connection", "NEW_UNLISTED": "New", "NO_CONNECTION": "No Connection"}
    for c in consumers:
        p = props.get(c.get("linked_property_id"), {})
        s = surveys.get(c.get("linked_property_id"), {})
        w = wards.get(c.get("ward_id"), {})
        docs = ", ".join(sorted({a["attachment_type"] for a in s.get("attachments", [])})) if s else ""
        water = [x["connection_number"] for x in c["connections"] if x["service"] == "Water"]
        sewer = [x["connection_number"] for x in c["connections"] if x["service"] == "Sewer"]
        row = [town.get("town") or town_code, w.get("ward_number") or c.get("ward_number") or p.get("ward") or ward_master.ward_for_colony(c.get("colony_name") or p.get("colony")) or "", c.get("colony_name") or p.get("colony") or "",
               p.get("property_id", ""), p.get("owner_name", ""), p.get("address", ""), p.get("latitude"), p.get("longitude"),
               c.get("consumer_id", ""), c.get("consumer_name", ""), c.get("fh_name", ""), c.get("ppp_id", ""), c.get("address", ""), c.get("locality", ""), c.get("phone", ""),
               ", ".join(water), ", ".join(sewer), len(water), len(sewer), len(water) + len(sewer), c.get("category", ""),
               type_label.get(s.get("survey_type"), "Existing" if c["source"] == "imported" else "New") if s or c else "",
               "Linked" if c.get("linked_property_id") else "Not linked", ("Yes" if s.get("owner_mismatch") else "No") if s else "",
               s.get("status", "Not Started"), s.get("surveyor_name", ""), s.get("latitude"), s.get("longitude"), s.get("gps_accuracy"),
               (s.get("submitted_at") or "")[:19].replace("T", " "), docs or ("None" if s else ""), s.get("remarks") or "",
               f"{c.get('source')}: {c.get('source_file') or ''} row {c.get('source_row') or ''}".strip()]
        ws.append(row)
        r = ws.max_row
        for col in (4, 9, 10, 11, 12, 15, 16, 17):
            ws.cell(row=r, column=col).number_format = "@"
    for col_cells in ws.columns:
        ws.column_dimensions[col_cells[0].column_letter].width = 18
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    await audit(user, "EXPORT", "export", "phed", meta={"rows": len(consumers)})
    fname = f"PHED_Survey_Export_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    return StreamingResponse(out, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": f'attachment; filename="{fname}"'})
