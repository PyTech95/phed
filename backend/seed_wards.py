"""Seed phed_wards (ward -> colonies) for Thanesar from the ward-wise colony list.
Idempotent by ward_number. Run: python seed_wards.py <xlsx_path>"""
import asyncio
import os
import sys
import uuid
from datetime import datetime, timezone
from openpyxl import load_workbook
from motor.motor_asyncio import AsyncIOMotorClient

MONGO_URL = os.environ.get("MONGO_URL", "mongodb://localhost:27017")
DB_NAME = os.environ.get("DB_NAME", "test_database")


def parse(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    header = [str(c).strip().lower() if c is not None else "" for c in rows[0]]
    col_i = next((i for i, h in enumerate(header) if "colony" in h), 2)
    ward_i = next((i for i, h in enumerate(header) if "ward" in h), 3)
    wards = {}
    for r in rows[1:]:
        if not r or len(r) <= max(col_i, ward_i):
            continue
        colony = r[col_i]
        ward = r[ward_i]
        if colony is None or ward is None:
            continue
        wn = str(ward).strip().replace("Ward", "").replace("ward", "").strip()
        try:
            wn = str(int(float(wn)))
        except Exception:
            pass
        wards.setdefault(wn, [])
        c = str(colony).strip()
        if c and c not in wards[wn]:
            wards[wn].append(c)
    return wards


async def main(path):
    wards = parse(path)
    client = AsyncIOMotorClient(MONGO_URL)
    db = client[DB_NAME]
    now = datetime.now(timezone.utc).isoformat()
    n = 0
    for wn, colonies in sorted(wards.items(), key=lambda x: (len(x[0]), x[0])):
        doc = {"ward_number": wn, "name": f"Ward {wn}", "town": "Thanesar",
               "is_active": True, "colonies": colonies, "updated_at": now}
        res = await db.phed_wards.update_one(
            {"ward_number": wn},
            {"$set": doc, "$setOnInsert": {"id": str(uuid.uuid4()), "created_at": now}},
            upsert=True)
        n += 1
        print(f"Ward {wn}: {len(colonies)} colonies {'(new)' if res.upserted_id else '(updated)'}")
    total = await db.phed_wards.count_documents({})
    print(f"DONE. wards processed={n}, total phed_wards={total}")


if __name__ == "__main__":
    asyncio.run(main(sys.argv[1] if len(sys.argv) > 1 else "/tmp/colony_list.xlsx"))
