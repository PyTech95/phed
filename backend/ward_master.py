"""Thanesar colony -> ward master, loaded from the bundled ward-wise colony list.
Provides colony->ward auto-mapping (import/field-create) and ward->colonies (seeding/dropdowns)."""
import logging
import re
from pathlib import Path

from openpyxl import load_workbook

_DATA = Path(__file__).parent / "data" / "colony_ward_thanesar.xlsx"
_colony_to_ward = {}
_ward_to_colonies = {}
_loaded = False


def _norm(s):
    return re.sub(r"\s+", " ", str(s or "").strip()).lower()


def _load():
    global _loaded
    if _loaded:
        return
    _loaded = True
    if not _DATA.exists():
        logging.getLogger("ward_master").warning("colony->ward master file not found: %s", _DATA)
        return
    try:
        wb = load_workbook(_DATA, read_only=True, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        header = [str(c).strip().lower() if c is not None else "" for c in rows[0]]
        col_i = next((i for i, h in enumerate(header) if "colony" in h), 2)
        ward_i = next((i for i, h in enumerate(header) if "ward" in h), 3)
        for r in rows[1:]:
            if not r or len(r) <= max(col_i, ward_i):
                continue
            colony, ward = r[col_i], r[ward_i]
            if colony is None or ward is None:
                continue
            wn = str(ward).strip()
            try:
                wn = str(int(float(wn)))
            except (ValueError, TypeError):
                pass
            c = str(colony).strip()
            if not c:
                continue
            _colony_to_ward[_norm(c)] = wn
            _ward_to_colonies.setdefault(wn, [])
            if c not in _ward_to_colonies[wn]:
                _ward_to_colonies[wn].append(c)
        logging.getLogger("ward_master").info("Loaded %d colonies across %d wards", len(_colony_to_ward), len(_ward_to_colonies))
    except Exception:
        logging.getLogger("ward_master").exception("Failed to load colony->ward master")


def ward_for_colony(colony):
    """Return the ward number (as string) for a colony name, or None if unknown."""
    _load()
    if not colony:
        return None
    return _colony_to_ward.get(_norm(colony))


def ward_to_colonies():
    """Return {ward_number: [colony names]} sorted by numeric ward."""
    _load()
    return {wn: list(cols) for wn, cols in sorted(_ward_to_colonies.items(), key=lambda x: (len(x[0]), x[0]))}
