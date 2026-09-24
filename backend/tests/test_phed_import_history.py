"""
Tests for PHED import history display + sample template download.
Bug context: importing already-existing consumers previously showed '+0 consumers'
so the UI looked like a failure. Backend must return updated_consumers count
and a downloadable .xlsx template.
"""
import os
import io
import pytest
import requests

BASE_URL = os.environ.get('REACT_APP_BACKEND_URL').rstrip('/')
ADMIN_USER = "admin"
ADMIN_PASS = "PhedAdmin@2026"
TOWN = "THS"


@pytest.fixture(scope="module")
def admin_token():
    r = requests.post(f"{BASE_URL}/api/auth/login",
                      json={"username": ADMIN_USER, "password": ADMIN_PASS},
                      timeout=30)
    assert r.status_code == 200, f"login failed: {r.status_code} {r.text}"
    return r.json()["token"]


@pytest.fixture(scope="module")
def h(admin_token):
    return {"Authorization": f"Bearer {admin_token}", "X-Town-Code": TOWN}


# --- Sample template download -------------------------------------------------
class TestImportSample:
    def test_sample_download_ok(self, h):
        r = requests.get(f"{BASE_URL}/api/phed/import/sample", headers=h, timeout=30)
        assert r.status_code == 200
        ct = r.headers.get("content-type", "")
        assert "openxmlformats-officedocument.spreadsheetml.sheet" in ct, ct
        assert len(r.content) > 500
        # verify workbook headers
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(r.content))
        ws = wb.active
        row1 = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        expected = [
            "Consumer Name", "F/H Name", "Head of Family in PPP", "Address",
            "Locality", "Phone No.", "Consumer ID", "Water Connection No.",
            "Sewer Connection No.", "Type of Connection",
        ]
        assert row1 == expected, row1

    def test_sample_requires_auth(self):
        r = requests.get(f"{BASE_URL}/api/phed/import/sample",
                         headers={"X-Town-Code": TOWN}, timeout=30)
        assert r.status_code in (401, 403)


# --- Import history contains updated_consumers -------------------------------
class TestImportHistory:
    def test_history_returns_updated_count(self, h):
        r = requests.get(f"{BASE_URL}/api/phed/imports", headers=h, timeout=60)
        assert r.status_code == 200
        data = r.json()
        items = data.get("imports") if isinstance(data, dict) else data
        assert isinstance(items, list) and len(items) > 0, f"no history: {data}"
        # find a DIDAR import
        didar = [i for i in items if "DIDAR" in (i.get("filename") or "").upper()]
        assert didar, f"no DIDAR entry in history; filenames={[i.get('filename') for i in items[:10]]}"
        latest = didar[0]
        res = latest.get("result") or {}
        assert "created_consumers" in res
        assert "updated_consumers" in res
        # For the reported file: 0 created, 3103 updated
        assert res["updated_consumers"] >= 1, f"expected updates, got {res}"

    def test_get_single_import_detail(self, h):
        r = requests.get(f"{BASE_URL}/api/phed/imports", headers=h, timeout=60)
        items = r.json().get("imports") if isinstance(r.json(), dict) else r.json()
        didar = [i for i in items if "DIDAR" in (i.get("filename") or "").upper()]
        assert didar
        iid = didar[0]["id"]
        r2 = requests.get(f"{BASE_URL}/api/phed/import/{iid}", headers=h, timeout=30)
        assert r2.status_code == 200
        d = r2.json()
        res = d.get("result") or {}
        assert res.get("created_consumers") == 0
        assert res.get("updated_consumers", 0) >= 3000, res


# --- Regression --------------------------------------------------------------
class TestRegression:
    def test_property_search_didar(self, h):
        r = requests.get(f"{BASE_URL}/api/admin/properties?search=Didar&limit=1",
                         headers=h, timeout=30)
        assert r.status_code == 200
        j = r.json()
        assert j.get("total", 0) >= 1000, j.get("total")

    def test_phed_consumer_search_ram(self, h):
        r = requests.get(f"{BASE_URL}/api/phed/consumers/search?q=RAM&limit=10",
                         headers=h, timeout=30)
        assert r.status_code == 200
        items = r.json()
        if isinstance(items, dict):
            items = items.get("items") or items.get("results") or []
        assert len(items) > 0
        # Prefix match preference: at least one starts with RAM
        names = [(i.get("consumer_name") or "").upper() for i in items]
        assert any(n.startswith("RAM") for n in names), names[:5]
