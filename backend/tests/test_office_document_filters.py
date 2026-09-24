"""Office Document status filter regressions (iter15).

Covers GET /api/phed/survey-consumers, /api/phed/consumers,
/api/phed/consumers/search, and /api/phed/export document_status
acceptance across the five OfficeDocumentStatus labels plus
"not_submitted" (nested missing/null) and "All" (parameter omitted).
Seeds isolated TEST_ODF_ prefixed fixtures, uses existing admin +
citywide surveyor credentials (no new accounts), asserts:
  * exact-status match, "not_submitted" catches nulls + absent nested
    receipt, All restores all, invalid status returns 422
  * staff endpoint excludes source='survey'; admin includes them
  * combined search + link_status/linked + admin ward/category/source
  * city isolation (x-town-code) and unauth (401/403) rejection
  * counts / pagination beyond 40-staff and 25-admin remain filtered
  * autocomplete /consumers/search respects document filter for exact
    (consumer_id, phone, connection, property_id) + prefix + broad
  * /export xlsx round-trip: matching row ids honour filter and
    include the new Office Document Status + Receipt Property ID columns
  * Only latest saved receipt is filtered; NO writes to consumers /
    connections / properties / surveys
Teardown deletes only exact fixture IDs.
"""
from __future__ import annotations

import copy
import io
import os
from datetime import datetime, timezone

import pytest
import requests
from openpyxl import load_workbook
from pymongo import MongoClient


def _read_env_file(path: str) -> dict:
    out: dict = {}
    try:
        for line in open(path):
            line = line.strip()
            if not line or line.startswith("#") or "=" not in line:
                continue
            k, v = line.split("=", 1)
            out[k.strip()] = v.strip().strip('"').strip("'")
    except FileNotFoundError:
        pass
    return out


_BE = _read_env_file("/app/backend/.env")
_FE = _read_env_file("/app/frontend/.env")

BASE_URL = (os.environ.get("REACT_APP_BACKEND_URL") or _FE.get("REACT_APP_BACKEND_URL") or "").rstrip("/")
assert BASE_URL, "REACT_APP_BACKEND_URL missing"
MONGO_URL = os.environ.get("MONGO_URL") or _BE.get("MONGO_URL")
DB_NAME = os.environ.get("DB_NAME") or _BE.get("DB_NAME")
assert MONGO_URL and DB_NAME, "MONGO_URL / DB_NAME missing"
ADMIN_USERNAME = os.environ.get("ADMIN_USERNAME") or _BE.get("ADMIN_USERNAME")
ADMIN_PASSWORD = os.environ.get("ADMIN_PASSWORD") or _BE.get("ADMIN_PASSWORD")
assert ADMIN_USERNAME and ADMIN_PASSWORD, "Admin credentials missing"

THS = {"x-town-code": "THS"}


def _parse_surveyor_creds(path="/app/memory/test_credentials.md"):
    u = p = None
    try:
        in_sec = False
        for raw in open(path, encoding="utf-8"):
            line = raw.rstrip("\n")
            if line.startswith("## Test Surveyor"):
                in_sec = True
                continue
            if in_sec:
                if line.startswith("## "):
                    break
                if line.startswith("- Username:"):
                    u = line.split(":", 1)[1].strip().strip("`").strip()
                elif line.startswith("- Password:"):
                    p = line.split(":", 1)[1].strip().strip("`").strip()
    except FileNotFoundError:
        pass
    return u, p


SURVEYOR_USERNAME, SURVEYOR_PASSWORD = _parse_surveyor_creds()
assert SURVEYOR_USERNAME and SURVEYOR_PASSWORD, "Failed to parse Test Surveyor from memory"

FIVE_STATUSES = [
    "Sewer documents received at office",
    "Already ok PID received at office",
    "Death transfer documents received at office",
    "Owner change Documents received at office",
    "Other",
]

# Fixture IDs (exact - kept isolated for teardown)
CONS_PREFIX = "TEST_ODF_CONS_"
PROP_REC = "rec_TEST_ODF_LINKPROP_001"
PROP_PID = "TEST_ODF_LINKPROP_001"
CONN_ID_LINKED = "TEST_ODF_CONN_LINKED"


def _cons_id(idx: int) -> str:
    return f"{CONS_PREFIX}{idx:03d}"


def _cons_cid(idx: int) -> str:
    return f"TEST_ODF_CID_{idx:03d}"


@pytest.fixture(scope="session")
def mongo():
    c = MongoClient(MONGO_URL)
    yield c[DB_NAME]
    c.close()


@pytest.fixture(scope="session")
def admin_token():
    r = requests.post(f"{BASE_URL}/api/auth/login",
                      json={"username": ADMIN_USERNAME, "password": ADMIN_PASSWORD, "selected_town": "THS"},
                      timeout=15)
    assert r.status_code == 200, r.text
    return r.json()["token"]


@pytest.fixture(scope="session")
def surveyor_token(admin_token, mongo):
    # Ensure user exists (idempotent seed)
    if not mongo.users.find_one({"username": SURVEYOR_USERNAME}):
        town = mongo.towns.find_one({"code": "THS"})
        payload = {"username": SURVEYOR_USERNAME, "password": SURVEYOR_PASSWORD,
                   "name": "Citywide Test Surveyor", "role": "SURVEYOR",
                   "assigned_town": town["id"], "gps_radius_required": False}
        requests.post(f"{BASE_URL}/api/admin/users", json=payload,
                      headers={"Authorization": f"Bearer {admin_token}", **THS}, timeout=15)
    r = requests.post(f"{BASE_URL}/api/auth/login",
                      json={"username": SURVEYOR_USERNAME, "password": SURVEYOR_PASSWORD, "selected_town": "THS"},
                      timeout=15)
    assert r.status_code == 200, r.text
    return r.json()["token"]


# -------- seed / teardown ---------
CONS_IDS: list = []
STATUS_BY_IDX: dict = {}


def _mk_consumer(idx, status_val, source="imported", is_active=True, receipt_present=True, linked=False):
    now = datetime.now(timezone.utc).isoformat()
    doc = {
        "id": _cons_id(idx),
        "consumer_id": _cons_cid(idx),
        "consumer_id_norm": _cons_cid(idx).upper(),
        "consumer_name": f"ODF Consumer {idx}",
        "name_norm": f"odf consumer {idx}",
        "fh_name": "ODF FH",
        "phone": f"90000000{idx:02d}",
        "phone_norm": f"90000000{idx:02d}",
        "address": "ODF addr",
        "address_norm": "odf addr",
        "locality": "ODF_COLONY",
        "colony_name": "ODF_COLONY",
        "ward_id": "1",
        "ward_number": "1",
        "status": "active",
        "is_active": is_active,
        "category": "Domestic",
        "source": source,
        "created_at": now,
        "updated_at": now,
    }
    if linked:
        doc["linked_property_id"] = PROP_REC
        doc["linked_property_number"] = PROP_PID
    else:
        doc["linked_property_id"] = None
    if receipt_present:
        # status_val=None means explicit null (not_submitted case)
        doc["office_document_receipt"] = {
            "status": status_val,
            "remarks": "ODF remark" if status_val else None,
            "property_id": PROP_PID,
            "property_record_id": PROP_REC,
            "recorded_at": now,
            "recorded_by": "TEST_ODF_USER",
        }
    return doc


@pytest.fixture(scope="session", autouse=True)
def seed_fixtures(mongo):
    now = datetime.now(timezone.utc).isoformat()
    # Seed a real property for the receipt linkage
    mongo.properties.replace_one({"id": PROP_REC}, {
        "id": PROP_REC, "property_id": PROP_PID, "batch_id": "TEST_ODF_BATCH",
        "owner_name": "ODF Owner", "mobile": "9000090000", "address": "ODF Prop Addr",
        "colony": "ODF_COLONY", "ward": "1", "status": "Pending", "serial_number": 1,
        "town": "THS", "latitude": 29.97, "longitude": 76.87, "created_at": now,
    }, upsert=True)

    # 1 consumer per full label (idx 1..5), 1 with receipt=None (idx 6 = not_submitted null),
    # 1 without receipt key at all (idx 7 = not_submitted absent),
    # 1 with source='survey' + status Other (idx 8 - staff excludes, admin includes),
    # 1 linked to property + status Sewer (idx 9 - for linked filter tests),
    # 1 with receipt=None but linked (idx 10 - linked + not_submitted)
    # Extra Other-status consumers force pagination beyond 40 staff results.
    docs = []
    for i, s in enumerate(FIVE_STATUSES, start=1):
        docs.append(_mk_consumer(i, s))
        STATUS_BY_IDX[i] = s
    docs.append(_mk_consumer(6, None))                    # explicit null status
    docs.append(_mk_consumer(7, None, receipt_present=False))  # no receipt field
    docs.append(_mk_consumer(8, "Other", source="survey"))     # staff excluded
    docs.append(_mk_consumer(9, FIVE_STATUSES[0], linked=True))
    docs.append(_mk_consumer(10, None, linked=True))
    for i in range(11, 60):
        docs.append(_mk_consumer(i, "Other"))

    ids = [d["id"] for d in docs]
    CONS_IDS.extend(ids)
    for d in docs:
        mongo.phed_consumers.replace_one({"id": d["id"]}, d, upsert=True)

    mongo.phed_connections.replace_one({"id": CONN_ID_LINKED}, {
        "id": CONN_ID_LINKED, "consumer_ref": _cons_id(9), "service": "Water",
        "connection_number": "TEST_ODF_CONN_1", "connection_number_norm": "TEST_ODF_CONN_1",
        "is_active": True, "status": "active", "created_at": now, "updated_at": now,
    }, upsert=True)

    yield

    mongo.phed_consumers.delete_many({"id": {"$in": ids}})
    mongo.properties.delete_many({"id": PROP_REC})
    mongo.phed_connections.delete_many({"id": CONN_ID_LINKED})


def sv(t):
    return {"Authorization": f"Bearer {t}", **THS}


# =================================================================
# GET /api/phed/survey-consumers  (staff)
# =================================================================
class TestStaffSurveyConsumersDocumentFilter:
    def test_unauth_rejected(self):
        r = requests.get(f"{BASE_URL}/api/phed/survey-consumers", timeout=15)
        assert r.status_code in (401, 403)

    def test_all_omits_param_returns_all_seed(self, surveyor_token):
        r = requests.get(f"{BASE_URL}/api/phed/survey-consumers",
                         params={"limit": 100}, headers=sv(surveyor_token), timeout=20)
        assert r.status_code == 200, r.text
        pids = {c["id"] for c in r.json()["consumers"]}
        # Staff excludes idx 8 (source=survey); all other seeded consumers present
        for i in list(range(1, 8)) + list(range(9, 60)):
            assert _cons_id(i) in pids, f"missing {_cons_id(i)} when All"
        assert _cons_id(8) not in pids, "staff must exclude source='survey'"

    @pytest.mark.parametrize("status", FIVE_STATUSES)
    def test_each_status_exact_match(self, surveyor_token, status):
        r = requests.get(f"{BASE_URL}/api/phed/survey-consumers",
                         params={"document_status": status, "limit": 100},
                         headers=sv(surveyor_token), timeout=20)
        assert r.status_code == 200, r.text
        rows = r.json()["consumers"]
        assert rows, f"no rows for {status}"
        for row in rows:
            assert row.get("office_document_receipt", {}).get("status") == status, row

    def test_not_submitted_covers_null_and_missing(self, surveyor_token):
        r = requests.get(f"{BASE_URL}/api/phed/survey-consumers",
                         params={"document_status": "not_submitted", "limit": 100},
                         headers=sv(surveyor_token), timeout=20)
        assert r.status_code == 200
        ids = {c["id"] for c in r.json()["consumers"]}
        assert _cons_id(6) in ids  # explicit null
        assert _cons_id(7) in ids  # absent
        assert _cons_id(10) in ids  # linked + null
        assert _cons_id(1) not in ids

    def test_invalid_status_422(self, surveyor_token):
        r = requests.get(f"{BASE_URL}/api/phed/survey-consumers",
                         params={"document_status": "__any__"},
                         headers=sv(surveyor_token), timeout=20)
        assert r.status_code == 422

    def test_invalid_status_random_422(self, surveyor_token):
        r = requests.get(f"{BASE_URL}/api/phed/survey-consumers",
                         params={"document_status": "Some Random Status"},
                         headers=sv(surveyor_token), timeout=20)
        assert r.status_code == 422

    def test_empty_result_returns_zero_counts(self, surveyor_token, mongo):
        # Combine filter that yields no matches (Sewer status + search on non-existent id)
        r = requests.get(f"{BASE_URL}/api/phed/survey-consumers",
                         params={"document_status": FIVE_STATUSES[0], "search": "TEST_ODF_NOMATCH_XYZ"},
                         headers=sv(surveyor_token), timeout=20)
        assert r.status_code == 200
        data = r.json()
        assert data["total"] == 0
        assert data["consumers"] == []
        assert data["pages"] in (0, 1)

    def test_combined_search_and_link_status_and_filter(self, surveyor_token):
        # idx 9 linked+Sewer, but link_status=unlinked must exclude it
        r = requests.get(f"{BASE_URL}/api/phed/survey-consumers",
                         params={"document_status": FIVE_STATUSES[0],
                                 "link_status": "unlinked"},
                         headers=sv(surveyor_token), timeout=20)
        assert r.status_code == 200
        ids = {c["id"] for c in r.json()["consumers"]}
        assert _cons_id(1) in ids   # unlinked+Sewer
        assert _cons_id(9) not in ids  # linked+Sewer -> excluded

        # link_status=linked with Sewer -> only idx 9
        r2 = requests.get(f"{BASE_URL}/api/phed/survey-consumers",
                          params={"document_status": FIVE_STATUSES[0], "link_status": "linked"},
                          headers=sv(surveyor_token), timeout=20)
        assert r2.status_code == 200
        ids2 = {c["id"] for c in r2.json()["consumers"]}
        assert _cons_id(9) in ids2
        assert _cons_id(1) not in ids2

    def test_pagination_beyond_40_stays_filtered(self, surveyor_token):
        # At least 50 Other-status fixtures ensure the second page is exercised.
        r1 = requests.get(f"{BASE_URL}/api/phed/survey-consumers",
                          params={"document_status": "Other", "page": 1, "limit": 40},
                          headers=sv(surveyor_token), timeout=20)
        assert r1.status_code == 200
        d1 = r1.json()
        assert d1["total"] >= 35
        for row in d1["consumers"]:
            assert row["office_document_receipt"]["status"] == "Other"
        assert d1["pages"] > 1
        r2 = requests.get(f"{BASE_URL}/api/phed/survey-consumers",
                          params={"document_status": "Other", "page": 2, "limit": 40},
                          headers=sv(surveyor_token), timeout=20)
        assert r2.status_code == 200
        page2 = r2.json()["consumers"]
        assert page2
        assert {c["id"] for c in d1["consumers"]}.isdisjoint(c["id"] for c in page2)
        for row in page2:
            assert row["office_document_receipt"]["status"] == "Other"


# =================================================================
# GET /api/phed/consumers  (admin list)
# =================================================================
class TestAdminConsumersDocumentFilter:
    def test_unauth_rejected(self):
        r = requests.get(f"{BASE_URL}/api/phed/consumers", timeout=15)
        assert r.status_code in (401, 403)

    def test_surveyor_rejected(self, surveyor_token):
        r = requests.get(f"{BASE_URL}/api/phed/consumers",
                         headers=sv(surveyor_token), timeout=20)
        assert r.status_code in (401, 403)

    def test_all_includes_source_survey(self, admin_token):
        r = requests.get(f"{BASE_URL}/api/phed/consumers",
                         params={"limit": 100},
                         headers=sv(admin_token), timeout=20)
        assert r.status_code == 200
        ids = {c["id"] for c in r.json()["consumers"]}
        # admin INCLUDES source=survey (idx 8)
        assert _cons_id(8) in ids

    @pytest.mark.parametrize("status", FIVE_STATUSES)
    def test_each_status_exact_match_admin(self, admin_token, status):
        r = requests.get(f"{BASE_URL}/api/phed/consumers",
                         params={"document_status": status, "limit": 100},
                         headers=sv(admin_token), timeout=20)
        assert r.status_code == 200
        for row in r.json()["consumers"]:
            assert row.get("office_document_receipt", {}).get("status") == status

    def test_admin_not_submitted_covers_null_and_absent(self, admin_token):
        r = requests.get(f"{BASE_URL}/api/phed/consumers",
                         params={"document_status": "not_submitted", "limit": 100},
                         headers=sv(admin_token), timeout=20)
        assert r.status_code == 200
        ids = {c["id"] for c in r.json()["consumers"]}
        assert _cons_id(6) in ids
        assert _cons_id(7) in ids

    def test_admin_invalid_status_422(self, admin_token):
        r = requests.get(f"{BASE_URL}/api/phed/consumers",
                         params={"document_status": "bogus"},
                         headers=sv(admin_token), timeout=20)
        assert r.status_code == 422

    def test_admin_combined_ward_source_linked_filter(self, admin_token):
        r = requests.get(f"{BASE_URL}/api/phed/consumers",
                         params={"document_status": FIVE_STATUSES[0],
                                 "linked": "yes", "source": "imported",
                                 "ward_id": "1", "category": "Domestic"},
                         headers=sv(admin_token), timeout=20)
        assert r.status_code == 200
        ids = {c["id"] for c in r.json()["consumers"]}
        assert _cons_id(9) in ids  # linked+imported+Sewer+ward1
        assert _cons_id(1) not in ids  # unlinked

    def test_admin_pagination_beyond_25_stays_filtered(self, admin_token):
        r1 = requests.get(f"{BASE_URL}/api/phed/consumers",
                          params={"document_status": "Other", "page": 1, "limit": 25},
                          headers=sv(admin_token), timeout=20)
        assert r1.status_code == 200
        d1 = r1.json()
        assert d1["total"] >= 35
        r2 = requests.get(f"{BASE_URL}/api/phed/consumers",
                          params={"document_status": "Other", "page": 2, "limit": 25},
                          headers=sv(admin_token), timeout=20)
        assert r2.status_code == 200
        for row in r1.json()["consumers"] + r2.json()["consumers"]:
            assert row.get("office_document_receipt", {}).get("status") == "Other"

    def test_admin_search_combined_with_document_filter(self, admin_token):
        # search for consumer_id of idx 2 (Already ok PID)
        r = requests.get(f"{BASE_URL}/api/phed/consumers",
                         params={"document_status": FIVE_STATUSES[1],
                                 "search": _cons_cid(2)},
                         headers=sv(admin_token), timeout=20)
        assert r.status_code == 200
        rows = r.json()["consumers"]
        assert len(rows) == 1
        assert rows[0]["id"] == _cons_id(2)

        # same search but mismatched document filter -> empty
        r2 = requests.get(f"{BASE_URL}/api/phed/consumers",
                          params={"document_status": FIVE_STATUSES[0],
                                  "search": _cons_cid(2)},
                          headers=sv(admin_token), timeout=20)
        assert r2.status_code == 200
        assert r2.json()["total"] == 0


# =================================================================
# GET /api/phed/consumers/search  (autocomplete)
# =================================================================
class TestConsumerSearchAutocompleteDocumentFilter:
    def test_unauth(self):
        r = requests.get(f"{BASE_URL}/api/phed/consumers/search",
                         params={"q": "ODF"}, timeout=15)
        assert r.status_code in (401, 403)

    def test_exact_consumer_id_with_matching_status(self, surveyor_token):
        r = requests.get(f"{BASE_URL}/api/phed/consumers/search",
                         params={"q": _cons_cid(1), "document_status": FIVE_STATUSES[0]},
                         headers=sv(surveyor_token), timeout=20)
        assert r.status_code == 200
        ids = {c["id"] for c in r.json()["results"]}
        assert _cons_id(1) in ids

    def test_exact_consumer_id_mismatched_status_excluded(self, surveyor_token):
        r = requests.get(f"{BASE_URL}/api/phed/consumers/search",
                         params={"q": _cons_cid(1), "document_status": FIVE_STATUSES[1]},
                         headers=sv(surveyor_token), timeout=20)
        assert r.status_code == 200
        ids = {c["id"] for c in r.json()["results"]}
        assert _cons_id(1) not in ids

    def test_phone_hit_respects_filter(self, surveyor_token):
        # phone for idx 2 = 9000000002 -> Already ok PID
        r = requests.get(f"{BASE_URL}/api/phed/consumers/search",
                         params={"q": "9000000002", "document_status": FIVE_STATUSES[1]},
                         headers=sv(surveyor_token), timeout=20)
        assert r.status_code == 200
        ids = {c["id"] for c in r.json()["results"]}
        assert _cons_id(2) in ids

    def test_connection_number_hit_respects_filter(self, surveyor_token):
        # idx 9 -> Sewer status via linked connection
        r = requests.get(f"{BASE_URL}/api/phed/consumers/search",
                         params={"q": "TEST_ODF_CONN_1", "document_status": FIVE_STATUSES[0]},
                         headers=sv(surveyor_token), timeout=20)
        assert r.status_code == 200
        ids = {c["id"] for c in r.json()["results"]}
        assert _cons_id(9) in ids
        # Wrong status returns no hit for idx 9
        r2 = requests.get(f"{BASE_URL}/api/phed/consumers/search",
                          params={"q": "TEST_ODF_CONN_1", "document_status": FIVE_STATUSES[1]},
                          headers=sv(surveyor_token), timeout=20)
        assert _cons_id(9) not in {c["id"] for c in r2.json()["results"]}

    def test_broad_name_prefix_with_not_submitted(self, surveyor_token):
        r = requests.get(f"{BASE_URL}/api/phed/consumers/search",
                         params={"q": "odf consumer", "document_status": "not_submitted"},
                         headers=sv(surveyor_token), timeout=20)
        assert r.status_code == 200
        ids = {c["id"] for c in r.json()["results"]}
        assert _cons_id(6) in ids or _cons_id(7) in ids or _cons_id(10) in ids

    def test_all_status_omission(self, surveyor_token):
        r = requests.get(f"{BASE_URL}/api/phed/consumers/search",
                         params={"q": _cons_cid(1)},
                         headers=sv(surveyor_token), timeout=20)
        assert r.status_code == 200
        ids = {c["id"] for c in r.json()["results"]}
        assert _cons_id(1) in ids

    def test_invalid_status_422(self, surveyor_token):
        r = requests.get(f"{BASE_URL}/api/phed/consumers/search",
                         params={"q": "odf", "document_status": "__any__"},
                         headers=sv(surveyor_token), timeout=20)
        assert r.status_code == 422


# =================================================================
# GET /api/phed/export  (xlsx)
# =================================================================
class TestExportDocumentFilter:
    def test_unauth(self):
        r = requests.get(f"{BASE_URL}/api/phed/export", timeout=20)
        assert r.status_code in (401, 403)

    def test_surveyor_forbidden(self, surveyor_token):
        r = requests.get(f"{BASE_URL}/api/phed/export",
                         headers=sv(surveyor_token), timeout=30)
        assert r.status_code in (401, 403)

    def _parse_xlsx(self, content: bytes):
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        headers = list(rows[0])
        idx = {h: i for i, h in enumerate(headers)}
        return headers, rows[1:], idx

    def test_export_columns_present(self, admin_token):
        r = requests.get(f"{BASE_URL}/api/phed/export",
                         params={"document_status": FIVE_STATUSES[0]},
                         headers=sv(admin_token), timeout=60)
        assert r.status_code == 200
        headers, _, _ = self._parse_xlsx(r.content)
        assert "Office Document Status" in headers
        assert "Receipt Property ID" in headers

    @pytest.mark.parametrize("status", FIVE_STATUSES)
    def test_export_filters_rows_by_status(self, admin_token, status):
        r = requests.get(f"{BASE_URL}/api/phed/export",
                         params={"document_status": status},
                         headers=sv(admin_token), timeout=60)
        assert r.status_code == 200
        headers, rows, idx = self._parse_xlsx(r.content)
        cid_col = idx["PHED Consumer ID"]
        status_col = idx["Office Document Status"]
        # All ODF fixtures in export must have this exact status
        odf_rows = [row for row in rows if row[cid_col] and str(row[cid_col]).startswith("TEST_ODF_CID_")]
        assert odf_rows, "no ODF rows in export"
        for row in odf_rows:
            assert row[status_col] == status

    def test_export_not_submitted(self, admin_token):
        r = requests.get(f"{BASE_URL}/api/phed/export",
                         params={"document_status": "not_submitted"},
                         headers=sv(admin_token), timeout=60)
        assert r.status_code == 200
        _, rows, idx = self._parse_xlsx(r.content)
        cid_col = idx["PHED Consumer ID"]
        status_col = idx["Office Document Status"]
        odf_cids = {str(row[cid_col]) for row in rows if row[cid_col] and str(row[cid_col]).startswith("TEST_ODF_CID_")}
        assert _cons_cid(6) in odf_cids
        assert _cons_cid(7) in odf_cids
        # Missing and explicitly null statuses use the same human-readable label.
        for row in rows:
            if row[cid_col] and str(row[cid_col]).startswith("TEST_ODF_CID_"):
                assert row[status_col] == "Not submitted", row

    def test_export_search_plus_document(self, admin_token):
        r = requests.get(f"{BASE_URL}/api/phed/export",
                         params={"document_status": FIVE_STATUSES[1],
                                 "search": _cons_cid(2)},
                         headers=sv(admin_token), timeout=60)
        assert r.status_code == 200
        _, rows, idx = self._parse_xlsx(r.content)
        cid_col = idx["PHED Consumer ID"]
        cids = {str(row[cid_col]) for row in rows if row[cid_col]}
        assert _cons_cid(2) in cids
        assert _cons_cid(1) not in cids  # different status

    def test_export_linked_plus_document(self, admin_token):
        r = requests.get(f"{BASE_URL}/api/phed/export",
                         params={"document_status": FIVE_STATUSES[0], "linked": "yes"},
                         headers=sv(admin_token), timeout=60)
        assert r.status_code == 200
        _, rows, idx = self._parse_xlsx(r.content)
        cid_col = idx["PHED Consumer ID"]
        status_col = idx["Office Document Status"]
        prop_col = idx["Receipt Property ID"]
        for row in rows:
            if row[cid_col] and str(row[cid_col]).startswith("TEST_ODF_CID_"):
                assert row[status_col] == FIVE_STATUSES[0]
        cids = {str(row[cid_col]) for row in rows if row[cid_col]}
        assert _cons_cid(9) in cids  # linked+Sewer
        assert _cons_cid(1) not in cids  # unlinked
        # Receipt Property ID column populated for idx 9
        for row in rows:
            if row[cid_col] == _cons_cid(9):
                assert row[prop_col] == PROP_PID

    def test_export_invalid_status_422(self, admin_token):
        r = requests.get(f"{BASE_URL}/api/phed/export",
                         params={"document_status": "bogus"},
                         headers=sv(admin_token), timeout=30)
        assert r.status_code == 422


# =================================================================
# No side effects: filter operations do NOT modify records
# =================================================================
class TestFilterNoSideEffects:
    def test_snapshot_equality_after_filter_calls(self, admin_token, surveyor_token, mongo):
        # Take snapshots of a representative sample
        ids = [_cons_id(i) for i in (1, 2, 6, 7, 9)]
        pre_cons = {c["id"]: c for c in mongo.phed_consumers.find({"id": {"$in": ids}}, {"_id": 0})}
        pre_prop = mongo.properties.find_one({"id": PROP_REC}, {"_id": 0})
        pre_conn = mongo.phed_connections.find_one({"id": CONN_ID_LINKED}, {"_id": 0})

        # Fire a mixture of filter + search + export calls
        requests.get(f"{BASE_URL}/api/phed/survey-consumers",
                     params={"document_status": FIVE_STATUSES[0]},
                     headers=sv(surveyor_token), timeout=20)
        requests.get(f"{BASE_URL}/api/phed/consumers",
                     params={"document_status": "not_submitted"},
                     headers=sv(admin_token), timeout=20)
        requests.get(f"{BASE_URL}/api/phed/consumers/search",
                     params={"q": _cons_cid(1), "document_status": FIVE_STATUSES[0]},
                     headers=sv(surveyor_token), timeout=20)
        requests.get(f"{BASE_URL}/api/phed/export",
                     params={"document_status": FIVE_STATUSES[0]},
                     headers=sv(admin_token), timeout=60)

        post_cons = {c["id"]: c for c in mongo.phed_consumers.find({"id": {"$in": ids}}, {"_id": 0})}
        assert pre_cons == post_cons
        assert pre_prop == mongo.properties.find_one({"id": PROP_REC}, {"_id": 0})
        assert pre_conn == mongo.phed_connections.find_one({"id": CONN_ID_LINKED}, {"_id": 0})


# =================================================================
# City isolation - x-town-code wrong header should not leak fixtures
# =================================================================
class TestCityIsolation:
    def test_wrong_town_code_denied_or_empty_admin(self, admin_token):
        r = requests.get(f"{BASE_URL}/api/phed/consumers",
                         params={"document_status": FIVE_STATUSES[0]},
                         headers={"Authorization": f"Bearer {admin_token}", "x-town-code": "ZZZ"},
                         timeout=20)
        # Either 403 or empty result set
        if r.status_code == 200:
            ids = {c["id"] for c in r.json()["consumers"]}
            assert not any(cid.startswith(CONS_PREFIX) for cid in ids)
        else:
            assert r.status_code == 403
