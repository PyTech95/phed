"""
Backend tests for the Colony->Ward master (iteration 8).
- GET /api/phed/wards seeded 32 wards / 192 colonies
- POST /api/phed/field-properties auto-derives ward from known colony
- POST /api/admin/batch/upload auto-derives ward from Colony header when Ward missing
- GET /api/admin/export contains Ward No column populated from master
- GET /api/phed/export Ward column populated from master fallback
"""
import io
import os
import uuid
import pytest
import requests
from openpyxl import Workbook, load_workbook

BASE_URL = (os.environ.get("REACT_APP_BACKEND_URL") or
            open("/app/frontend/.env").read().split("REACT_APP_BACKEND_URL=")[1].splitlines()[0].strip()
            ).rstrip("/")
TOWN = "THS"
HB = {"X-Town-Code": TOWN}


def _login(username, password):
    r = requests.post(f"{BASE_URL}/api/auth/login",
                      json={"username": username, "password": password, "selected_town": TOWN},
                      headers={**HB, "Content-Type": "application/json"}, timeout=30)
    assert r.status_code == 200, r.text
    return r.json().get("access_token") or r.json().get("token")


@pytest.fixture(scope="module")
def admin_tok():
    return _login(os.environ.get("TEST_ADMIN_USERNAME", "admin"), os.environ.get("TEST_ADMIN_PASSWORD", "PhedAdmin@2026"))


@pytest.fixture(scope="module")
def surveyor_tok():
    return _login("surveyor1", "Survey@2026")


@pytest.fixture(scope="module")
def admin_h(admin_tok):
    return {**HB, "Authorization": f"Bearer {admin_tok}", "Content-Type": "application/json"}


@pytest.fixture(scope="module")
def surveyor_h(surveyor_tok):
    return {**HB, "Authorization": f"Bearer {surveyor_tok}", "Content-Type": "application/json"}


# ---------- 1. Ward master seeded ----------
def test_wards_seeded(surveyor_h):
    r = requests.get(f"{BASE_URL}/api/phed/wards", headers=surveyor_h, timeout=30)
    assert r.status_code == 200, r.text
    j = r.json()
    wards = j.get("wards") or []
    assert len(wards) == 32, f"expected 32 wards, got {len(wards)}"
    total = sum(len(w.get("colonies") or []) for w in wards)
    assert total == 192, f"expected 192 total colonies, got {total}"

    by_num = {str(w.get("ward_number")): (w.get("colonies") or []) for w in wards}
    # samples
    assert any("kirti nagar" == c.lower() for c in by_num.get("4", [])), by_num.get("4")
    assert any("saraswati colony jhansa road" == c.lower() for c in by_num.get("1", [])), by_num.get("1")
    # no empty wards
    for w in wards:
        assert w.get("colonies"), f"ward {w.get('ward_number')} has no colonies"


# ---------- 2. Field property auto-derives ward from colony ----------
def test_field_property_ward_auto(surveyor_h):
    body = {"owner_name": f"TEST_WardAuto_{uuid.uuid4().hex[:6]}",
            "mobile": "9998887777", "ward": "", "colony": "Kirti Nagar",
            "category": "Residential", "latitude": 29.9695, "longitude": 76.8783}
    r = requests.post(f"{BASE_URL}/api/phed/field-properties", json=body,
                      headers=surveyor_h, timeout=30)
    assert r.status_code == 200, r.text
    prop = r.json()["property"]
    assert prop["ward"] == "4", f"expected ward '4' auto-derived, got {prop.get('ward')!r}"


def test_field_property_ward_unknown_colony_no_error(surveyor_h):
    body = {"owner_name": f"TEST_UnknownColony_{uuid.uuid4().hex[:6]}",
            "mobile": "9998887778", "ward": "", "colony": "ZZZ_UnknownColony",
            "category": "Residential", "latitude": 29.97, "longitude": 76.88}
    r = requests.post(f"{BASE_URL}/api/phed/field-properties", json=body,
                      headers=surveyor_h, timeout=30)
    assert r.status_code == 200, r.text
    prop = r.json()["property"]
    # ward should be empty or fall back to colony (no crash)
    assert prop["ward"] in ("", "ZZZ_UnknownColony"), prop.get("ward")


# ---------- 3. Batch upload: ward auto-derived from Colony when no Ward col ----------
def _build_xlsx(rows, headers):
    wb = Workbook()
    ws = wb.active
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h, "") for h in headers])
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out


def _find_property(admin_h, property_id):
    r = requests.get(f"{BASE_URL}/api/admin/properties", params={"limit": 1000}, headers=admin_h, timeout=60)
    assert r.status_code == 200, r.text
    data = r.json()
    props = data.get("properties") if isinstance(data, dict) else data
    return next((p for p in props if p.get("property_id") == property_id), None)


def test_batch_upload_ward_auto_from_colony(admin_tok, admin_h):
    pid = f"TEST-WA-{uuid.uuid4().hex[:6].upper()}"
    xlsx = _build_xlsx(
        [{"Property Id": pid, "Owner Name": "TEST_BatchWard", "Colony": "Kirti Nagar"}],
        ["Property Id", "Owner Name", "Colony"],
    )
    files = {"file": ("wardtest.xlsx", xlsx, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    data = {"batch_name": "WardTest", "authorization": admin_tok}
    r = requests.post(f"{BASE_URL}/api/admin/batch/upload", files=files, data=data,
                      headers=HB, timeout=60)
    assert r.status_code == 200, r.text
    prop = _find_property(admin_h, pid)
    assert prop is not None, f"uploaded property {pid} not found"
    assert prop.get("ward") == "4", f"expected ward '4' auto-derived, got {prop.get('ward')!r}"


def test_batch_upload_explicit_ward_wins(admin_tok, admin_h):
    pid = f"TEST-WE-{uuid.uuid4().hex[:6].upper()}"
    xlsx = _build_xlsx(
        [{"Property Id": pid, "Owner Name": "TEST_ExplicitWard", "Colony": "Kirti Nagar", "Ward No": "99"}],
        ["Property Id", "Owner Name", "Colony", "Ward No"],
    )
    files = {"file": ("wardtest2.xlsx", xlsx, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    data = {"batch_name": "WardTest2", "authorization": admin_tok}
    r = requests.post(f"{BASE_URL}/api/admin/batch/upload", files=files, data=data,
                      headers=HB, timeout=60)
    assert r.status_code == 200, r.text
    prop = _find_property(admin_h, pid)
    assert prop is not None
    assert prop.get("ward") == "99", f"explicit ward should win, got {prop.get('ward')!r}"


# ---------- 4. Admin export contains Ward No column populated ----------
def test_admin_export_has_ward_column(admin_h):
    r = requests.get(f"{BASE_URL}/api/admin/export?status=", headers=admin_h, timeout=120)
    assert r.status_code == 200, r.status_code
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb.active
    headers = [c.value for c in ws[1]]
    assert "Ward No" in headers, f"'Ward No' missing from headers: {headers}"
    ward_col = headers.index("Ward No") + 1
    colony_col = headers.index("Colony Name") + 1
    # find at least one row whose colony maps in master and confirm Ward No filled
    found = False
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True):
        colony = str(row[colony_col - 1] or "").strip().lower()
        ward = str(row[ward_col - 1] or "").strip()
        if colony == "kirti nagar" and ward == "4":
            found = True
            break
    assert found, "No exported row with Kirti Nagar -> ward 4 found"


# ---------- 5. PHED export Ward column populated ----------
def test_phed_export_has_ward(admin_h):
    r = requests.get(f"{BASE_URL}/api/phed/export", headers=admin_h, timeout=120)
    assert r.status_code == 200, r.status_code
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb.active
    headers = [c.value for c in ws[1]]
    assert "Ward" in headers, f"'Ward' missing from headers: {headers}"
    # Column exists; if rows present, at least ensure header ok.
    # Additional: if any row has a colony known to master, ward should be non-empty
    if ws.max_row >= 2:
        colony_col = headers.index("Colony") + 1
        ward_col = headers.index("Ward") + 1
        for row in ws.iter_rows(min_row=2, max_row=min(ws.max_row, 500), values_only=True):
            colony = str(row[colony_col - 1] or "").strip().lower()
            if colony == "kirti nagar":
                assert str(row[ward_col - 1] or "").strip() == "4"
                return
